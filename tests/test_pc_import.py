"""
회귀 테스트 — scripts/pc_import.py · scripts/import_sources.py (이 PC 내역 정리)

실행:
    python3 -m unittest discover tests            (웹·리눅스)
    $PY -m unittest discover tests                (이 PC — PY = D:/household/.venv/Scripts/python.exe)

내역 파일·CSV·규칙은 전부 시험 안에서 만든 합성 값이다(가짜 상호·가짜 금액). 저장소에는 텍스트만 둔다 —
xlsx/xls는 시험 중 임시 폴더에 합성 데이터로 만든다(탐색 기준선 3-0 11). 실제 거래내역을 넣지 않는다.

지키는 것(내역 자동화 탐색 기준선 4절 완료 기준 1~23·29 · 6절 K1~K13 · 시험 후보 1~14).
"""
import datetime as dt
import hashlib
import io
import json
import os
import re
import shutil
import sqlite3
import stat
import struct
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from unittest import mock

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "scripts"))

import import_sources as srcmod  # noqa: E402
import pc_import  # noqa: E402

TODAY = dt.date(2026, 10, 27)
NOW = dt.datetime(2026, 10, 27, 12, 0, 0, tzinfo=pc_import.KST)

CONFIG = r'''
start_date = "2026-10-01"
date_slack = { card = 2, bank = 0 }
platform_keys = ["가짜페이"]

[[source]]
id = "현대"
kind = "card"
card = "현대카드"
account = "생활비"
file_types = ["text", "xlsx", "xls"]
columns = ["이용일자", "이용시각", "가맹점명", "이용금액", "할부개월", "승인번호", "상태"]
identity = ["1111"]
date_basis = "승인일"
date_formats = ["%Y.%m.%d"]
installment_mode = "total"
cancel_values = ["취소"]
pay_day = 12
cycle_months_before = 1
cycle_start_day = 1
footer_re = "^합계"
period_re = '(\d{4})\.(\d{2})\.(\d{2})\s*~\s*(\d{4})\.(\d{2})\.(\d{2})'
[source.roles]
date = "이용일자"
time = "이용시각"
merchant = "가맹점명"
amount = "이용금액"
installment_months = "할부개월"
approval = "승인번호"
cancel = "상태"
[source.values]
"이용일자" = { type = "s", re = '\d{4}\.\d{2}\.\d{2}' }
"이용시각" = { type = "s", re = '\d{2}:\d{2}:\d{2}' }
"가맹점명" = { type = "s" }
"이용금액" = { type = ["s", "n"], re = '-?[\d,]+' }
"할부개월" = { type = "s", re = '일시불|\d+개월' }
"승인번호" = { type = "s", re = '\d{8}' }
"상태" = { type = "s", re = '정상|취소' }

[[source]]
id = "신한카드"
kind = "card"
card = "신한카드"
account = "생활비"
file_types = ["text", "xlsx", "xls"]
columns = ["거래일", "가맹점", "금액", "승인번호"]
identity = ["2222"]
date_basis = "승인일"
date_formats = ["%Y-%m-%d"]
pay_day = 15
[source.roles]
date = "거래일"
merchant = "가맹점"
amount = "금액"
approval = "승인번호"
[source.values]
"거래일" = { type = ["s", "d"], re = '\d{4}-\d{2}-\d{2}( 00:00:00)?' }
"가맹점" = { type = "s" }
"금액" = { type = ["s", "n"], re = '-?[\d,]+' }
"승인번호" = { type = "s", re = '\d{8}' }

[[source]]
id = "생활비"
kind = "bank"
account = "생활비"
file_types = ["text", "xlsx", "xls"]
columns = ["거래일시", "적요", "내용", "출금", "입금", "잔액"]
identity = ["3333-01"]
date_formats = ["%Y-%m-%d %H:%M:%S"]
[source.roles]
date = "거래일시"
merchant = "내용"
memo = ["적요", "내용"]
out = "출금"
in = "입금"
balance = "잔액"
[source.values]
"거래일시" = { type = "s", re = '\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}' }
"적요" = { type = "s" }
"내용" = { type = "s", empty_ok = true }
"출금" = { type = ["s", "n"], re = '[\d,]+' }
"입금" = { type = ["s", "n"], re = '[\d,]+' }
"잔액" = { type = ["s", "n"], re = '-?[\d,]+' }

[[source]]
id = "신한은행"
kind = "bank"
account = "신한은행"
file_types = ["text", "xlsx", "xls"]
columns = ["거래일자", "거래시간", "적요", "출금(원)", "입금(원)", "내용", "잔액(원)"]
identity = ["110-222"]
date_formats = ["%Y.%m.%d"]
[source.roles]
date = "거래일자"
time = "거래시간"
merchant = "내용"
memo = ["적요", "내용"]
out = "출금(원)"
in = "입금(원)"
balance = "잔액(원)"
[source.values]
"거래일자" = { type = "s", re = '\d{4}\.\d{2}\.\d{2}' }
"거래시간" = { type = "s", re = '\d{2}:\d{2}:\d{2}' }
"적요" = { type = "s" }
"출금(원)" = { type = "s", re = '[\d,]+' }
"입금(원)" = { type = "s", re = '[\d,]+' }
"내용" = { type = "s", empty_ok = true }
"잔액(원)" = { type = "s", re = '-?[\d,]+' }

[[transfer]]
account = "생활비"
counterpart = "신한은행"
patterns = ['^신한 110222$']

[[transfer]]
account = "생활비"
counterpart = "청년미래적금"
patterns = ["적금통장"]

[[transfer]]
account = "생활비"
counterpart = "커플통장"
patterns = ["커플모임통장"]

[[transfer]]
account = "생활비"
counterpart = "주식잔고"
patterns = ["주식계좌"]

[[transfer]]
account = "신한은행"
counterpart = "생활비"
patterns = ['^카카오 3333-01$']

[[transfer_ambiguous]]
account = "생활비"
patterns = ["가나다"]
candidates = ["커플통장", "신한은행"]
default = "커플통장"

[[transfer_pair]]
from = "생활비"
to = "신한은행"
"대분류" = "이체"
"세부내용" = "고정비 이체"
"고정여부" = "고정"

[[transfer_pair]]
from = "생활비"
to = "청년미래적금"
"대분류" = "이체"
"세부내용" = "적금 이체"
"고정여부" = "고정"

[[transfer_pair]]
from = "생활비"
to = "커플통장"
"대분류" = "이체"
"세부내용" = "커플통장 이체"
"고정여부" = "고정"

[[card_payment]]
account = "생활비"
card = "현대카드"
patterns = ['^현대카드$']

[[card_payment]]
account = "생활비"
card = "신한카드"
patterns = ['^신한카드$']

[[income]]
account = "생활비"
patterns = ['^가상회사$']
"대분류" = "급여"
"소분류" = ""
"세부내용" = "월급"
"고정여부" = "고정"
'''

RULES = (
    "키,대분류,소분류,세부내용,고정여부,결제수단,물음,원문,만든 날,근거 회차\r\n"
    "가상마트,식비,장보기,가상마트,,,자동,가상마트 2호점,2026-10-01,20261001-000000\r\n"
    "가상구독,문화생활비,어플/멤버쉽,가상구독,고정,신한은행,자동,가상구독,2026-10-01,20261001-000000\r\n"
)

HEADER = "날짜,구분,대분류,소분류,세부내용, 금액 ,결제수단,고정여부,비고"
BASE_ROWS = [
    '2026-09-01,수입,급여,,9월 급여,"3,123,457",생활비,고정,',
    '2026-09-01,이체,이체,신한은행,고정비 이체,"701,234",생활비,고정,',
    '2026-09-01,이체,이체,청년미래적금,적금 이체,"501,234",생활비,고정,',
    '2026-09-01,이체,이체,커플통장,커플통장 이체,"401,234",생활비,고정,',
    '2026-09-05,지출,식비,외식,"점심, 회사 근처","12,345",생활비,,현대카드',
    '2026-09-10,지출,식비,장보기,마트,48123,생활비,,신한카드',
    '2026-09-15,지출,교통비,대중교통,버스,517,생활비,,',
    '2026-10-01,수입, 급여 ,,10월 급여,"3,123,457",생활비,고정,',
    '2026-10-01,이체,이체,신한은행,고정비 이체,"701,234",생활비,고정,',
    '2026-10-01,이체,이체,청년미래적금,적금 이체,"501,234",생활비,고정,',
    '2026-10-01,이체,이체,커플통장,커플통장 이체,"401,234",생활비,고정,',
    '2026-10-02,지출,식비,장보기,마트,"8,812",생활비,,신한카드',
    '2026-10-03,지출,식비,외식, 점심 ,"15,013",생활비,,현대카드',
    '2026-09-20,지출,문화생활비,어플/멤버쉽,가상구독,"9,917",신한은행,고정,현대카드',
]


def csv_bytes(rows=BASE_ROWS, tail=True, bom=True, nl="\r\n"):
    text = nl.join([HEADER] + list(rows) + ([",,,,,,,,"] if tail else [])) + nl
    return (b"\xef\xbb\xbf" if bom else b"") + text.encode("utf-8")


def md5(b):
    return hashlib.md5(b).hexdigest()


def read_bytes(path):
    with open(path, "rb") as f:
        return f.read()


def _rm_readonly(func, path, _):
    os.chmod(path, stat.S_IWRITE)
    func(path)


def rmtree(path):
    if sys.version_info >= (3, 12):
        shutil.rmtree(path, onexc=_rm_readonly)
    else:
        shutil.rmtree(path, onerror=_rm_readonly)


# ---- 합성 내역 파일 ----

def text_file(head, columns, rows):
    lines = [",".join(h) for h in head] + [",".join(columns)]
    for r in rows:
        lines.append(",".join(f'"{c}"' if "," in str(c) else str(c) for c in r))
    return ("\r\n".join(lines) + "\r\n").encode("utf-8")


HY_COLS = ["이용일자", "이용시각", "가맹점명", "이용금액", "할부개월", "승인번호", "상태"]
SH_COLS = ["거래일", "가맹점", "금액", "승인번호"]
LV_COLS = ["거래일시", "적요", "내용", "출금", "입금", "잔액"]
SB_COLS = ["거래일자", "거래시간", "적요", "출금(원)", "입금(원)", "내용", "잔액(원)"]


def hyundai(rows, ident="1111", period=None):
    """rows: (날짜 'MM-DD', 시각, 가맹점, 금액, 할부, 승인번호, 상태)"""
    head = [["현대카드 이용내역"], ["카드번호", f"****-****-****-{ident}"]]
    if period:
        head.append(["조회기간", f"{period[0]} ~ {period[1]}"])
    head.append([])
    data = [[f"2026.{d.replace('-', '.')}", t, m, f"{a:,}", inst, ap, st] for d, t, m, a, inst, ap, st in rows]
    return text_file(head, HY_COLS, data)


def shinhan_card(rows, ident="2222"):
    head = [["신한카드", f"****{ident}"]]
    data = [[f"2026-{d}", m, f"{a:,}", ap] for d, m, a, ap in rows]
    return text_file(head, SH_COLS, data)


def living(rows, start=1_000_013, ident="3333-01", newest_first=False):
    """rows: (날짜 'MM-DD', 시각, 적요, 내용, 출금, 입금) — 잔액은 계산한다."""
    bal, data = start, []
    for d, t, memo, content, o, i in rows:
        bal = bal - o + i
        data.append([f"2026-{d} {t}", memo, content, f"{o:,}", f"{i:,}", f"{bal:,}"])
    if newest_first:
        data.reverse()
    return text_file([["카카오뱅크 거래내역"], ["계좌번호", ident]], LV_COLS, data)


def shinhan_bank(rows, start=50_123, ident="110-222-333337", newest_first=True):
    bal, data = start, []
    for d, t, memo, content, o, i in rows:
        bal = bal - o + i
        data.append([f"2026.{d.replace('-', '.')}", t, memo, f"{o:,}", f"{i:,}", content, f"{bal:,}"])
    if newest_first:
        data.reverse()
    return text_file([["신한은행 거래내역조회"], ["계좌", ident]], SB_COLS, data)


def write_xlsx(path, rows, sheet="Sheet1"):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(r)
    wb.save(path)


def _rec(op, data):
    return struct.pack("<HH", op, len(data)) + data


def _bof(kind):
    return _rec(0x0809, struct.pack("<HHHHII", 0x0600, kind, 0x0DBB, 0x07CC, 0, 0x06))


def _ustr(s, lenlen):
    return struct.pack("<B" if lenlen == 1 else "<H", len(s)) + b"\x01" + s.encode("utf-16-le")


def _xf(fmt):
    return _rec(0x00E0, struct.pack("<HHHBBBBIiH", 0, fmt, 0x0001, 0x20, 0, 0, 0, 0, 0, 0x20C0))


def write_xls(path, rows, sheet="Sheet1"):
    """최소 BIFF8 스트림(OLE 상자 없음 — xlrd가 그대로 읽는다). 문자열·숫자·날짜(XF 14) 칸만."""
    cells = b""
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            if v is None or v == "":
                continue
            if isinstance(v, dt.datetime):
                serial = (v - dt.datetime(1899, 12, 30)).total_seconds() / 86400
                cells += _rec(0x0203, struct.pack("<HHHd", r, c, 1, serial))
            elif isinstance(v, (int, float)):
                cells += _rec(0x0203, struct.pack("<HHHd", r, c, 0, float(v)))
            else:
                cells += _rec(0x0204, struct.pack("<HHH", r, c, 0) + _ustr(str(v), 2))
    body = _bof(0x0010) + cells + _rec(0x000A, b"")

    def glob(offset):
        return _bof(0x0005) + _xf(0) + _xf(14) + _rec(0x0085, struct.pack("<iBB", offset, 0, 0) + _ustr(sheet, 1)) \
            + _rec(0x000A, b"")
    with open(path, "wb") as f:
        f.write(glob(len(glob(0))) + body)


class World:
    """임시 작업 루트: uploads/가계부T.csv · uploads/내역/<날짜>/ · rules/(계좌.toml·분류규칙.csv)."""

    def __init__(self, tc, rows=BASE_ROWS, config=CONFIG, rules=RULES):
        self.tc = tc
        self.home = tempfile.mkdtemp(prefix="pcimport-")
        tc.addCleanup(rmtree, self.home)
        self.uploads = os.path.join(self.home, "uploads")
        self.inbox = os.path.join(self.uploads, "내역")
        self.rules = os.path.join(self.home, "rules")
        os.makedirs(self.inbox)
        os.makedirs(self.rules)
        with open(os.path.join(self.uploads, "가계부T.csv"), "wb") as f:
            f.write(csv_bytes(rows) if isinstance(rows, list) else rows)
        with open(os.path.join(self.rules, "계좌.toml"), "w", encoding="utf-8") as f:
            f.write(config)
        if rules is not None:
            with open(os.path.join(self.rules, "분류규칙.csv"), "wb") as f:
                f.write(b"\xef\xbb\xbf" + rules.encode("utf-8"))
        self.clock = [NOW]

    def put(self, name, data, folder="2026-10-27"):
        d = os.path.join(self.inbox, folder)
        os.makedirs(d, exist_ok=True)
        path = os.path.join(d, name)
        if callable(data):
            data(path)
        else:
            with open(path, "wb") as f:
                f.write(data)
        return path

    def csv_path(self):
        names = [n for n in os.listdir(self.uploads) if n.lower().endswith(".csv") and not n.startswith("~$")]
        self.tc.assertEqual(len(names), 1, names)
        return os.path.join(self.uploads, names[0])

    def csv(self):
        return read_bytes(self.csv_path())

    def run(self, *argv, today=TODAY):
        buf = io.StringIO()
        self.clock[0] = self.clock[0] + dt.timedelta(seconds=1)
        now = dt.datetime(today.year, today.month, today.day, self.clock[0].hour, self.clock[0].minute,
                          self.clock[0].second, tzinfo=pc_import.KST)
        with mock.patch.object(pc_import, "today_kst", lambda: today), \
                mock.patch.object(pc_import, "now_kst", lambda: now), redirect_stdout(buf):
            rc = pc_import.main(["--home", self.home, "--repo", ROOT, *argv])
        return rc, buf.getvalue()

    def scan(self, today=TODAY):
        rc, out = self.run("scan", today=today)
        runs = pc_import.run_dirs(pc_import.Paths(ROOT, self.home, False))
        return rc, out, (runs[-1] if runs else None)

    def run_dir(self, rid):
        return os.path.join(self.home, "work", "내역", rid)

    def items(self, rid):
        import openpyxl
        wb = openpyxl.load_workbook(os.path.join(self.run_dir(rid), "확인필요.xlsx"))
        ws = wb["확인 필요"]
        rows = list(ws.iter_rows(values_only=True))
        head = rows[0]
        return [dict(zip(head, r)) for r in rows[1:]]

    def answer(self, rid, pick, **cols):
        """pick(항목 dict) → True인 줄의 답 칸을 채운다."""
        import openpyxl
        path = os.path.join(self.run_dir(rid), "확인필요.xlsx")
        wb = openpyxl.load_workbook(path)
        ws = wb["확인 필요"]
        head = [c.value for c in ws[1]]
        n = 0
        for row in ws.iter_rows(min_row=2):
            d = {h: c.value for h, c in zip(head, row)}
            if pick(d):
                for k, v in cols.items():
                    row[head.index(k)].value = v
                n += 1
        wb.save(path)
        self.tc.assertGreater(n, 0, "답할 항목이 없다")
        return n

    def apply(self, rid, *extra, confirm=True, today=TODAY):
        return self.run("apply", "--run", rid, *extra, *(["--confirm"] if confirm else []), today=today)

    def ledger(self, sql):
        con = sqlite3.connect(os.path.join(self.rules, "원장.sqlite"))
        try:
            return con.execute(sql).fetchall()
        finally:
            con.close()

    def proposals(self, rid):
        import csv as _csv
        with open(os.path.join(self.run_dir(rid), "제안.csv"), encoding="utf-8-sig", newline="") as f:
            return list(_csv.DictReader(f))

    def new_lines(self, before):
        """새 CSV에서 옛 CSV 뒤에 덧붙은 데이터 줄(끝 빈 줄 제외)."""
        old = before.decode("utf-8-sig").split("\r\n")
        new = self.csv().decode("utf-8-sig").split("\r\n")
        tail = [x for x in old if x in (",,,,,,,,", "")]
        body_old = [x for x in old if x not in tail]
        return [x for x in new if x not in body_old and x not in (",,,,,,,,", "")]


ACK = ["--ack-balance", "생활비", "시험 기준 차액", "--ack-balance", "신한은행", "시험 기준 차액",
       "--ack-warning", "transfer_issues", "시험: 이체 합계"]
NEW_MONTH = ["--ack-warning", "transfer_issues", "시험: 새 달", "--ack-warning", "card_zero", "시험: 새 달"]


def cancel_ack(*cards):
    out = []
    for c in cards:
        out += ["--ack-card", c, "시험: 이용기간 파일 없음"]
    return out


class Base(unittest.TestCase):
    def world(self, **kw):
        return World(self, **kw)

    def ok_scan(self, w, today=TODAY):
        rc, out, rid = w.scan(today=today)
        self.assertIn(rc, (0, 2), out)
        return out, rid


# ============================================================
# 파서 · 서명 (완료 기준 2 · 15 · K11 · K12 · 시험 후보 7·9)
# ============================================================

class SourceTests(Base):
    def cfg(self):
        path = os.path.join(tempfile.mkdtemp(prefix="pcimport-cfg-"), "c.toml")
        self.addCleanup(rmtree, os.path.dirname(path))
        with open(path, "w", encoding="utf-8") as f:
            f.write(CONFIG)
        import pipeline
        return srcmod.load_config(path, set(pipeline.ACCOUNT_KEY.values()), set(pipeline.CARD_TARGETS))

    def read(self, path, sid):
        cfg = self.cfg()
        src = next(s for s in cfg["sources"] if s.id == sid)
        kind, sheets = srcmod.read_table(path)
        hit = srcmod.identify(sheets, src)
        self.assertIsNotNone(hit)
        txs = srcmod.parse_sheet(sheets[hit[0]][1], hit[1], src, os.path.basename(path))
        srcmod.check_balance_chain(txs, path)
        return srcmod.assign_keys(txs, src)

    def tmpfile(self, name, data=None, writer=None):
        d = tempfile.mkdtemp(prefix="pcimport-f-")
        self.addCleanup(rmtree, d)
        path = os.path.join(d, name)
        if writer:
            writer(path)
        else:
            with open(path, "wb") as f:
                f.write(data)
        return path

    def test_role_values_text(self):
        """완료 기준 15: 자동 행 값 = 역할 열 값(금액 정수 · 방향 · 승인일)."""
        path = self.tmpfile("h.csv", hyundai([("10-05", "19:00:00", "가상마트 2호점", 30123, "일시불", "10000002", "정상"),
                                              ("10-06", "08:00:00", "가상가게", 1234567, "3개월", "10000003", "정상")]))
        a, b = self.read(path, "현대")
        self.assertEqual((a.date, a.time, a.amount, a.direction, a.merchant, a.approval, a.inst_months),
                         (dt.date(2026, 10, 5), "19:00:00", 30123, "spend", "가상마트 2호점", "10000002", 0))
        self.assertEqual((b.amount, b.inst_months), (1234567, 3))

    def test_xlsx_and_xls_cells(self):
        """셀 형(datetime·숫자) 서명 — 엑셀에서 다시 저장해 형이 바뀐 파일도 서명에 적은 형이면 읽는다."""
        rows = [["신한카드", "****2222"], SH_COLS,
                [dt.datetime(2026, 10, 4), "가상마트 3호점", 48000, "20000001"]]
        for name, writer in (("s.xlsx", write_xlsx), ("s.xls", write_xls)):
            path = self.tmpfile(name, writer=lambda p, w=writer: w(p, rows))
            (t,) = self.read(path, "신한카드")
            self.assertEqual((t.date, t.amount, t.merchant), (dt.date(2026, 10, 4), 48000, "가상마트 3호점"), name)

    def test_value_signature_mismatch_fails(self):
        """시험 후보 7: 열 이름 같고 값 꼴이 다르면(MM.DD · 2026/10/05 · '1,234원' · 형 다름) 그 출처 실패."""
        cfg = self.cfg()
        src = next(s for s in cfg["sources"] if s.id == "현대")
        bad_rows = [
            ["10.05", "19:00:00", "가상가게", "1,013", "일시불", "10000009", "정상"],
            ["2026/10/05", "19:00:00", "가상가게", "1,013", "일시불", "10000009", "정상"],
            ["2026.10.05", "19:00:00", "가상가게", "1,234원", "일시불", "10000009", "정상"],
            ["2026.10.05", "19:00:00", "가상가게", "1013", "일시불", "1000013", "정상"],
        ]
        for row in bad_rows:
            data = text_file([["카드번호", "****1111"]], HY_COLS, [row])
            path = self.tmpfile("bad.csv", data)
            kind, sheets = srcmod.read_table(path)
            hit = srcmod.identify(sheets, src)
            with self.assertRaises(srcmod.SourceError, msg=row):
                srcmod.parse_sheet(sheets[hit[0]][1], hit[1], src, "bad.csv")
        # 셀 형: 숫자만 허용한 열에 문자열 → 실패
        cfg2 = self.cfg()
        lv = next(s for s in cfg2["sources"] if s.id == "생활비")
        lv.values["출금"].types = {"n"}
        path = self.tmpfile("lv.csv", living([("10-01", "09:00:00", "이체", "x", 1000, 0)]))
        kind, sheets = srcmod.read_table(path)
        hit = srcmod.identify(sheets, lv)
        with self.assertRaises(srcmod.SourceError):
            srcmod.parse_sheet(sheets[hit[0]][1], hit[1], lv, "lv.csv")

    def test_broken_files_fail(self):
        """시험 후보 9: HTML 본문 .xls(로그인 화면) · 암호 xlsx · 0바이트 → 실패(건너뜀 아님)."""
        w = self.world()
        w.put("login.xls", b"<html><body><table><tr><td>login</td></tr></table></body></html>")
        out, rid = self.ok_scan(w)
        self.assertIn("어느 출처 서명", out)
        rc, out = w.apply(rid, confirm=False)
        self.assertEqual(rc, 1, out)
        cases = {
            "empty.xlsx": b"",
            "enc.xlsx": b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64 + "EncryptedPackage".encode("utf-16-le"),
        }
        for name, data in cases.items():
            path = self.tmpfile(name, data)
            with self.assertRaises(srcmod.SourceError, msg=name):
                srcmod.read_table(path)

    def test_identity_mismatch_fails(self):
        """K12: 다른 카드·계좌 파일(식별 값 다름)은 서명 통과가 아니라 실패."""
        cfg = self.cfg()
        src = next(s for s in cfg["sources"] if s.id == "현대")
        path = self.tmpfile("h.csv", hyundai([("10-05", "19:00:00", "가상가게", 1013, "일시불", "10000002", "정상")],
                                             ident="9999"))
        kind, sheets = srcmod.read_table(path)
        with self.assertRaises(srcmod.SourceError):
            srcmod.identify(sheets, src)

    def test_balance_chain(self):
        """3-0 8(a): 잔액이 이어지지 않는 파일(행이 빠짐)은 실패 · 최신순 파일도 읽는다."""
        rows = [("10-01", "09:00:00", "급여", "가상회사", 0, 3123457), ("10-02", "10:00:00", "체크", "가상가게", 5013, 0)]
        path = self.tmpfile("lv.csv", living(rows, newest_first=True))
        txs = self.read(path, "생활비")
        self.assertEqual(len(txs), 2)
        data = living(rows).decode("utf-8").replace("4,118,457", "4,113,457").encode("utf-8")
        path = self.tmpfile("lv2.csv", data)
        with self.assertRaises(srcmod.SourceError):
            self.read(path, "생활비")

    def test_table_cut_fails(self):
        """빈 줄 뒤에 거래 행이 더 있는 파일(표가 끊김)은 실패."""
        data = hyundai([("10-05", "19:00:00", "가상가게", 1013, "일시불", "10000002", "정상")])
        data += b"\r\n2026.10.06,19:00:00,\xea\xb0\x80,1000,\xec\x9d\xbc\xec\x8b\x9c\xeb\xb6\x88,10000003,\xec\xa0\x95\xec\x83\x81\r\n"
        path = self.tmpfile("h.csv", data)
        with self.assertRaises(srcmod.SourceError):
            self.read(path, "현대")

    def test_normalize_branch_only(self):
        """완료 기준 12: 정규화는 끝의 지점 접미만 뗀다(가상마트24≠가상마트, 가상25≠가상)."""
        n = srcmod.normalize_merchant
        self.assertEqual(n("가상마트 2호점"), "가상마트")
        self.assertEqual(n("가상마트12점"), "가상마트")
        self.assertEqual(n("가상마트24"), "가상마트24")
        self.assertEqual(n("가상25 역삼점"), "가상25역삼점")
        self.assertEqual(n("ABC Shop"), "abcshop")

    def test_key_sequence(self):
        """완료 기준 4 · K9: 거래키에 같은 값 순번 — 같은 날 같은 금액 2건 → 키 2개. 최신순·오래된순이 같은 키."""
        cfg = self.cfg()
        src = next(s for s in cfg["sources"] if s.id == "신한은행")
        rows = [("10-02", "", "체크", "가상가게", 5013, 0), ("10-02", "", "체크", "가상가게", 5013, 0)]
        cfg_text = CONFIG.replace('time = "거래시간"\n', "")
        data = shinhan_bank([("10-02", "11:00:00", "체크", "가상가게", 5013, 0),
                             ("10-02", "11:00:00", "체크", "가상가게", 5013, 0)])
        path = self.tmpfile("sb.csv", data)
        txs = self.read(path, "신한은행")
        self.assertEqual(len({t.key for t in txs}), 2)
        data2 = shinhan_bank([("10-02", "11:00:00", "체크", "가상가게", 5013, 0),
                              ("10-02", "11:00:00", "체크", "가상가게", 5013, 0)], newest_first=False)
        path2 = self.tmpfile("sb2.csv", data2)
        self.assertEqual({t.key for t in txs}, {t.key for t in self.read(path2, "신한은행")})
        del src, rows, cfg_text

    def test_config_requires_signatures(self):
        """서명 없는 열·식별 값 없음·키 약함 미표시 → 설정 오류(파싱하지 않는다)."""
        import pipeline
        bad = CONFIG.replace('"상태" = { type = "s", re = \'정상|취소\' }\n', "")
        bad = bad.replace('identity = ["2222"]\n', "")
        path = self.tmpfile("c.toml", bad.encode("utf-8"))
        with self.assertRaises(srcmod.ConfigError) as cm:
            srcmod.load_config(path, set(pipeline.ACCOUNT_KEY.values()), set(pipeline.CARD_TARGETS))
        msg = "\n".join(cm.exception.args[0])
        self.assertIn("상태", msg)
        self.assertIn("identity", msg)


# ============================================================
# scan — 쓰기 0 · 짝 · 분류 (완료 기준 1·3·5·6·11·12·21)
# ============================================================

def basic_files(w):
    w.put("hyundai.csv", hyundai([
        ("10-03", "12:30:00", "가상식당 1호점", 15013, "일시불", "10000001", "정상"),   # 손 행과 짝(같은 날)
        ("10-05", "19:00:00", "가상마트 2호점", 30123, "일시불", "10000002", "정상"),   # 규칙 자동
        ("10-06", "08:00:00", "새가게", 7013, "일시불", "10000003", "정상"),            # 새 가맹점
        ("10-07", "09:00:00", "가짜페이 결제", 20013, "일시불", "10000004", "정상"),     # 플랫폼
    ]))
    w.put("living.csv", living([
        ("10-01", "09:00:00", "급여", "가상회사", 0, 3_123_457),
        ("10-01", "10:00:00", "타행이체", "신한 110222", 701_234, 0),
        ("10-01", "10:05:00", "이체", "적금통장", 501_234, 0),
        ("10-01", "10:10:00", "이체", "커플모임통장", 401_234, 0),
        ("10-15", "11:00:00", "입금", "가나다", 0, 50_123),
        ("10-20", "10:00:00", "이체", "주식계좌", 100_123, 0),
    ]))
    w.put("shinhan_bank.csv", shinhan_bank([
        ("10-01", "10:00:01", "타행입금", "카카오 3333-01", 0, 701_234),
    ]))


class ScanTests(Base):
    def test_scan_writes_nothing_outside_run(self):
        """완료 기준 1: scan은 uploads·rules·원장·log·clone 전후 md5 같음. 쓰기는 회차 폴더뿐."""
        w = self.world()
        basic_files(w)

        def tree(*roots):
            out = {}
            for root in roots:
                for r, _, fs in os.walk(root):
                    for n in fs:
                        out[os.path.join(r, n)] = md5(read_bytes(os.path.join(r, n)))
            return out
        before = tree(w.uploads, w.rules)
        clone_before = tree(os.path.join(ROOT, "scripts"), os.path.join(ROOT, "tests"))
        out, rid = self.ok_scan(w)
        self.assertEqual(before, tree(w.uploads, w.rules), out)
        self.assertEqual(clone_before, tree(os.path.join(ROOT, "scripts"), os.path.join(ROOT, "tests")))
        self.assertFalse(os.path.exists(os.path.join(w.home, "runs", "log.md")))
        for name in ("stamp.json", "확인필요.xlsx", "제안.csv", "보고.md", "거래.json", f"초안_가계부_{rid}.csv"):
            self.assertTrue(os.path.exists(os.path.join(w.run_dir(rid), name)), name)

    def test_basic_dispositions(self):
        """규칙 자동 · 손 행 짝 · 새 가맹점/플랫폼은 확인 필요 · 이체 한 행 · 제외 계좌 이체 자동 · 본인 이름 이체는 자동 0."""
        w = self.world()
        basic_files(w)
        out, rid = self.ok_scan(w)
        props = {(p["출처"], p["날짜"], p["금액"]): p for p in w.proposals(rid)}
        self.assertEqual(props[("현대", "2026-10-03", "15013")]["상태"], "짝", out)
        self.assertEqual(props[("현대", "2026-10-05", "30123")]["상태"], "자동")
        self.assertEqual(props[("현대", "2026-10-06", "7013")]["상태"], "확인 필요")
        self.assertEqual(props[("현대", "2026-10-07", "20013")]["상태"], "확인 필요")
        self.assertEqual(props[("생활비", "2026-10-01", "3123457")]["상태"], "짝")       # 급여 손 행
        self.assertEqual(props[("생활비", "2026-10-01", "701234")]["상태"], "짝")        # 이체 손 행(양쪽 내역)
        self.assertEqual(props[("신한은행", "2026-10-01", "701234")]["상태"], "제외")    # 입금 쪽 = 짝
        self.assertEqual(props[("생활비", "2026-10-01", "501234")]["상태"], "짝")        # 제외 계좌 이체 — 손 행 있음
        self.assertEqual(props[("생활비", "2026-10-20", "100123")]["상태"], "자동")      # 제외 계좌 이체 — 손 행 없음
        self.assertEqual(props[("생활비", "2026-10-15", "50123")]["상태"], "확인 필요")  # 본인 이름 → 자동 0
        kinds = {i["종류"] for i in w.items(rid)}
        self.assertTrue({"new_merchant", "platform", "transfer_unknown"} <= kinds, kinds)

    def test_rescan_same_files_no_new(self):
        """완료 기준 3: 같은 파일 두 번·기간 겹침 → apply 뒤 새 행 0."""
        w = self.world()
        basic_files(w)
        out, rid = self.ok_scan(w)
        rc, out = w.apply(rid, *ACK)
        self.assertEqual(rc, 0, out)
        # 같은 파일을 다른 날 폴더에 한 번 더(기간 겹침)
        w.put("hyundai_again.csv", read_bytes(os.path.join(w.inbox, "2026-10-27", "hyundai.csv")), folder="2026-10-28")
        before = w.csv()
        out, rid2 = self.ok_scan(w)
        draft = read_bytes(os.path.join(w.run_dir(rid2), f"초안_가계부_{rid2}.csv"))
        self.assertEqual(draft, before, out)

    def test_start_date_filter(self):
        """완료 기준 3 · K13: 시작일 전 거래는 버린다(1년 창 파일 두 회차에도 옛 달 추가 0)."""
        w = self.world()
        rows = [(f"0{m}-10", "10:00:00", f"옛가게{m}", 1000 * m, "일시불", f"3123457{m}", "정상") for m in range(1, 10)]
        w.put("old.csv", hyundai(rows, period=("2026.01.01", "2026.10.27")))
        out, rid = self.ok_scan(w)
        self.assertEqual([p for p in w.proposals(rid) if p["출처"] == "현대"], [], out)
        self.assertIn("시작일 전(버림) 9", out)

    def test_missing_file_is_skipped_failure_blocks(self):
        """완료 기준 21: 파일 없음 = 건너뜀(막지 않음) · 실패 = apply 막힘 → --skip(실패에만)으로만 연다."""
        w = self.world()
        w.put("hyundai.csv", hyundai([("10-05", "19:00:00", "가상마트 2호점", 30123, "일시불", "10000002", "정상")]))
        out, rid = self.ok_scan(w)
        self.assertIn("건너뜀(파일 없음)", out)
        self.assertNotIn("출처 실패", out)
        w.put("bad.csv", b"\xea\xb0\x80,\xeb\x82\x98\r\n1,2\r\n")
        out, rid = self.ok_scan(w)
        self.assertIn("어느 출처 서명", out)
        rc, out = w.apply(rid, confirm=False)
        self.assertEqual(rc, 1, out)
        rc, out = w.apply(rid, "--skip", "현대", confirm=False)
        self.assertEqual(rc, 1, out)
        self.assertIn("--skip은 '실패'에만", out)
        rc, out = w.apply(rid, "--skip", "2026-10-27/bad.csv", confirm=False)
        self.assertEqual(rc, 0, out)

    def test_ignore_file_needs_md5(self):
        """암호 원본 같은 파일은 설정(경로+md5)으로만 무시 — md5가 다르면 다시 실패."""
        enc = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64 + "EncryptedPackage".encode("utf-16-le")
        cfg = CONFIG + f'\n[[ignore_files]]\npath = "2026-10-27/enc.xlsx"\nmd5 = "{md5(enc)}"\nreason = "암호 원본"\n'
        w = self.world(config=cfg)
        w.put("enc.xlsx", enc)
        out, rid = self.ok_scan(w)
        self.assertIn("무시(설정)", out)
        self.assertNotIn("출처 실패", out)
        w.put("enc.xlsx", enc + b"x")
        out, rid = self.ok_scan(w)
        self.assertIn("출처 실패", out)


# ============================================================
# 손 행 짝 (완료 기준 5 · 6 · K2 · K5 · 시험 후보 10)
# ============================================================

class PairTests(Base):
    def test_one_of_two_same_amount_hand_rows(self):
        """같은 금액 2건 중 하나만 손으로 적었고 날짜가 다르면: 거리 0은 짝, 나머지는 자동 추가가 아니라 확인 필요."""
        rows = BASE_ROWS + ['2026-10-10,지출,식비,장보기,가상마트,"30,123",생활비,,현대카드']
        w = self.world(rows=rows)
        w.put("h.csv", hyundai([("10-09", "10:00:00", "가상마트 2호점", 30123, "일시불", "10000011", "정상"),
                                ("10-11", "10:00:00", "가상마트 2호점", 30123, "일시불", "10000012", "정상")]))
        out, rid = self.ok_scan(w)
        st = [p["상태"] for p in w.proposals(rid) if p["출처"] == "현대" and p["금액"] == "30123"]
        self.assertEqual(sorted(st), ["확인 필요", "확인 필요"], out)
        self.assertEqual(read_bytes(os.path.join(w.run_dir(rid), f"초안_가계부_{rid}.csv")), w.csv())

    def test_gap_pair_needs_answer(self):
        """D>0 짝은 자동 확정이 아니라 '확인 필요(짝 — 날짜 차)' — 답 '짝'이면 행 0."""
        rows = BASE_ROWS + ['2026-10-11,지출,식비,장보기,가상마트,"30,123",생활비,,현대카드']
        w = self.world(rows=rows)
        w.put("h.csv", hyundai([("10-09", "10:00:00", "가상마트 2호점", 30123, "일시불", "10000011", "정상")]))
        out, rid = self.ok_scan(w)
        items = [i for i in w.items(rid) if i["종류"] == "pair_gap"]
        self.assertEqual(len(items), 1, out)
        w.answer(rid, lambda d: d["종류"] == "pair_gap", 처리="짝")
        before = w.csv()
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        self.assertEqual(w.new_lines(before), [])

    def test_route_differs_and_future_typo(self):
        """비고 빠진 손 행(경로 다름) · 미래 날짜 손 행(날짜 오타) → 자동 추가 0."""
        rows = BASE_ROWS + ['2026-10-09,지출,식비,장보기,가상마트,"30,123",생활비,,',
                            '2026-11-12,지출,식비,외식,가상식당,"8,013",생활비,,현대카드']
        w = self.world(rows=rows)
        w.put("h.csv", hyundai([("10-09", "10:00:00", "가상마트 2호점", 30123, "일시불", "10000011", "정상"),
                                ("10-12", "10:00:00", "가상마트 2호점", 8013, "일시불", "10000012", "정상")]))
        out, rid = self.ok_scan(w)
        kinds = sorted(i["종류"] for i in w.items(rid))
        self.assertIn("pair_route", kinds, out)
        self.assertIn("pair_future", kinds, out)
        self.assertEqual(read_bytes(os.path.join(w.run_dir(rid), f"초안_가계부_{rid}.csv")), w.csv())

    def test_override_payment_card_row_vs_hand_row_without_memo(self):
        """검증 B2: 규칙이 결제수단을 바꾼 카드 행(가상구독 → 신한은행)과 비고 빠진 손 행(신한은행) →
        경로 다른 짝 '확인 필요' · 자동 추가 0(apply 뒤 새 행 0)."""
        rows = BASE_ROWS + ['2026-10-20,지출,문화생활비,어플/멤버쉽,가상구독,"9,917",신한은행,고정,']
        w = self.world(rows=rows)
        w.put("h.csv", hyundai([("10-20", "10:00:00", "가상구독", 9917, "일시불", "10000031", "정상")]))
        out, rid = self.ok_scan(w)
        props = [p for p in w.proposals(rid) if p.get("상태") == "자동"]
        self.assertEqual(props, [], out)
        self.assertIn("pair_route", [i["종류"] for i in w.items(rid)], out)
        self.assertEqual(read_bytes(os.path.join(w.run_dir(rid), f"초안_가계부_{rid}.csv")), w.csv())
        before = w.csv()
        rc, out = w.apply(rid, *ACK)
        self.assertEqual(rc, 0, out)
        self.assertEqual(w.new_lines(before), [], out)

    def test_csv_only_explained_not_asked_again(self):
        """완료 기준 6: 파일 기간 안인데 내역에 없는 CSV 행 → 답 '현금' → 다음 회차엔 묻지 않는다."""
        w = self.world()
        w.put("h.csv", hyundai([("10-05", "19:00:00", "가상마트 2호점", 30123, "일시불", "10000002", "정상")],
                               period=("2026.10.01", "2026.10.27")))
        out, rid = self.ok_scan(w)
        only = [i for i in w.items(rid) if i["종류"] == "csv_only"]
        self.assertEqual(len(only), 1, out)          # 10-03 현대카드 손 행(파일에 없음)
        w.answer(rid, lambda d: d["종류"] == "csv_only", 처리="현금")
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        out, rid2 = self.ok_scan(w)
        self.assertEqual([i for i in w.items(rid2) if i["종류"] == "csv_only"], [], out)

    def test_dropped_rows_do_not_shift_keys(self):
        """K5: load_csv가 버리는 행(빈 날짜·못 읽는 날짜)이 있어도 행 키는 값 기반 — 지울 줄은 정확히 그 줄."""
        doc_rows = BASE_ROWS[:5] + [',지출,식비,외식,날짜없음,"1,000",생활비,,현대카드'] + BASE_ROWS[5:]
        w = self.world(rows=doc_rows)
        doc = pc_import.CsvDoc(w.csv_path(), TODAY)
        keys = [r.key for r in doc.recs if r.key]
        self.assertEqual(len(keys), len(set(keys)))
        row = next(r for r in doc.recs if r.v["세부내용"] == "점심" and r.date == dt.date(2026, 10, 3))
        data, _ = doc.build([], {row.key})
        text = data.decode("utf-8-sig")
        self.assertNotIn("10-03,지출,식비,외식, 점심 ", text)
        self.assertIn("날짜없음", text)


# ============================================================
# CSV 쓰기 (완료 기준 16 · 17 · 29 · 시험 후보 5·6)
# ============================================================

class WriteTests(Base):
    def test_old_lines_byte_identical_new_rows_appended(self):
        """완료 기준 16: 옛 줄 바이트 그대로(같은 순서) · 새 행은 끝 빈 줄 앞에 한 가지 꼴 · BOM·CRLF 유지."""
        w = self.world()
        w.put("h.csv", hyundai([("10-05", "19:00:00", "가상마트 2호점", 30123, "일시불", "10000002", "정상"),
                                ("10-04", "19:00:00", "가상마트 2호점", 519, "일시불", "10000005", "정상"),
                                ("10-03", "12:30:00", "가상식당 1호점", 15013, "일시불", "10000001", "정상")]))
        before = w.csv()
        out, rid = self.ok_scan(w)
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        after = w.csv()
        self.assertTrue(after.startswith(b"\xef\xbb\xbf"))
        old_lines = before.decode("utf-8-sig").split("\r\n")
        new_lines = after.decode("utf-8-sig").split("\r\n")
        self.assertEqual(new_lines[: len(old_lines) - 2], old_lines[:-2])          # 끝 빈 줄 앞까지 그대로
        self.assertEqual(new_lines[len(old_lines) - 2: len(old_lines)],
                         ["2026-10-04,지출,식비,장보기,가상마트,519,생활비,,현대카드",
                          '2026-10-05,지출,식비,장보기,가상마트,"30,123",생활비,,현대카드'])
        self.assertEqual(new_lines[-2:], [",,,,,,,,", ""])
        self.assertNotIn(b"\n", after.replace(b"\r\n", b""))

    def test_one_csv_in_uploads_and_backup(self):
        """완료 기준 29: apply 뒤 uploads에 CSV 한 파일(새 이름) · 옛 파일은 backup/<회차>/ · 백업 md5 = 옛 것."""
        w = self.world()
        w.put("h.csv", hyundai([("10-05", "19:00:00", "가상마트 2호점", 30123, "일시불", "10000002", "정상")]))
        before = w.csv()
        out, rid = self.ok_scan(w)
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        self.assertEqual(os.path.basename(w.csv_path()), f"가계부_{rid}.csv")
        self.assertEqual(md5(read_bytes(os.path.join(w.uploads, "backup", rid, "가계부T.csv"))), md5(before))
        log = open(os.path.join(w.home, "runs", "log.md"), encoding="utf-8").read()
        self.assertIn(f"내역 {rid}", log)
        self.assertNotIn("30,123", log)
        self.assertNotIn("30123", log)
        self.assertNotIn("가상마트", log)

    def test_header_order_and_mismatch(self):
        """시험 후보 5: 열 순서 섞인 CSV → 헤더 순서대로 쓴다 · 열 집합이 다르면 멈춤."""
        hdr = "구분,날짜,대분류,소분류,세부내용,금액,결제수단,고정여부,비고"
        rows = ["수입,2026-10-01,급여,,10월 급여,\"3,123,457\",생활비,고정,", "지출,2026-10-02,식비,외식,x,1013,생활비,,"]
        data = ("\r\n".join([hdr] + rows) + "\r\n").encode("utf-8")
        w = self.world(rows=data)
        w.put("h.csv", hyundai([("10-05", "19:00:00", "가상마트 2호점", 30123, "일시불", "10000002", "정상")]))
        out, rid = self.ok_scan(w)
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        self.assertIn('지출,2026-10-05,식비,장보기,가상마트,"30,123",생활비,,현대카드', w.csv().decode("utf-8"))
        w2 = self.world(rows=data.replace(b"\xeb\xb9\x84\xea\xb3\xa0", b"memo"))   # 비고 → memo
        w2.put("h.csv", hyundai([("10-05", "19:00:00", "가상마트 2호점", 30123, "일시불", "10000002", "정상")]))
        rc, out, rid = w2.scan()
        self.assertEqual(rc, 1, out)
        self.assertIn("머리 9열", out)

    def test_warning_increase_blocks(self):
        """완료 기준 17: 쓰기 전 검사 — 경고 종류별 건수가 늘면(예: 미지 결제수단) 교체 0, --ack-warning으로만."""
        rules = RULES + "음수가게,식비,외식,음수가게,,,자동,음수가게,2026-10-01,x\r\n"
        w = self.world(rules=rules)
        # 결제수단 override가 계좌명이 아닌 규칙 → 새 행이 unknown_payment 경고를 늘린다
        rules_bad = RULES + "이상가게,식비,외식,이상가게,,엉뚱계좌,자동,이상가게,2026-10-01,x\r\n"
        with open(os.path.join(w.rules, "분류규칙.csv"), "wb") as f:
            f.write(b"\xef\xbb\xbf" + rules_bad.encode("utf-8"))
        w.put("h.csv", hyundai([("10-05", "19:00:00", "이상가게", 30123, "일시불", "10000002", "정상")]))
        before = w.csv()
        out, rid = self.ok_scan(w)
        rc, out = w.apply(rid)
        self.assertEqual(rc, 1, out)
        self.assertIn("쓰기 전 검사 accounts", out)
        self.assertEqual(w.csv(), before)

    def test_locked_csv_stops(self):
        """시험 후보 4: 엑셀이 CSV를 잡고 있으면(잠금 파일) 쓰기 0으로 멈춘다."""
        w = self.world()
        w.put("h.csv", hyundai([("10-05", "19:00:00", "가상마트 2호점", 30123, "일시불", "10000002", "정상")]))
        out, rid = self.ok_scan(w)
        lock = os.path.join(w.uploads, "~$가계부T.csv")
        with open(lock, "wb") as f:
            f.write(b"x")
        before = w.csv()
        rc, out = w.apply(rid)
        os.remove(lock)
        self.assertEqual(rc, 1, out)
        self.assertIn("엑셀", out)
        self.assertEqual(w.csv(), before)
        self.assertFalse(os.path.exists(os.path.join(w.uploads, "backup")))

    def test_apply_requires_latest_run_and_same_inputs(self):
        """완료 기준 19: 가장 최근 scan 회차만 · scan 뒤 입력이 바뀌면 apply 0."""
        w = self.world()
        w.put("h.csv", hyundai([("10-05", "19:00:00", "가상마트 2호점", 30123, "일시불", "10000002", "정상")]))
        out, rid1 = self.ok_scan(w)
        out, rid2 = self.ok_scan(w)
        rc, out = w.apply(rid1)
        self.assertEqual(rc, 1, out)
        self.assertIn("가장 최근", out)
        w.put("more.csv", shinhan_card([("10-08", "가상마트 5호점", 1000, "20000009")]))
        rc, out = w.apply(rid2)
        self.assertEqual(rc, 1, out)
        self.assertIn("입력이 바뀌었다", out)

    def test_crash_after_csv_written_relinks(self):
        """공통 안전망: 새 CSV를 쓴 뒤 원장 전에 죽어도 다음 scan이 새 행을 손 행처럼 짝지어 중복 0."""
        w = self.world()
        w.put("h.csv", hyundai([("10-05", "19:00:00", "가상마트 2호점", 30123, "일시불", "10000002", "정상")]))
        out, rid = self.ok_scan(w)
        with mock.patch.object(pc_import.sqlite3, "connect", side_effect=RuntimeError("시험: 죽음")):
            with self.assertRaises(RuntimeError):
                w.apply(rid)
        self.assertEqual(json.load(open(os.path.join(w.run_dir(rid), "apply.json"), encoding="utf-8"))["phase"], "rules")
        after = w.csv()
        out, rid2 = self.ok_scan(w)
        self.assertIn("지난 apply", out)
        self.assertEqual(read_bytes(os.path.join(w.run_dir(rid2), f"초안_가계부_{rid2}.csv")), after)
        props = [p for p in w.proposals(rid2) if p["출처"] == "현대"]
        self.assertEqual([p["상태"] for p in props], ["짝"])


# ============================================================
# 이체 (완료 기준 7 · K1 · 시험 후보 11)
# ============================================================

class TransferTests(Base):
    def test_two_rounds_in_then_out(self):
        """입금 파일 먼저(출금 계좌 파일 없음) → 출금 파일 다음 회차: 행 1."""
        rows = BASE_ROWS
        w = self.world(rows=rows)
        w.put("sb.csv", shinhan_bank([("10-02", "10:00:01", "타행입금", "카카오 3333-01", 0, 100_123)]))
        out, rid = self.ok_scan(w)
        # 생활비 파일이 없으니 입금 쪽은 '출금 계좌?' 확인 필요(자동 0)
        self.assertEqual([i["종류"] for i in w.items(rid) if i["출처"] == "신한은행"], ["transfer_noout"], out)
        w.answer(rid, lambda d: d["종류"] == "transfer_noout", 처리="추가")
        before = w.csv()
        rc, out = w.apply(rid, *ACK)
        self.assertEqual(rc, 0, out)
        self.assertEqual(len(w.new_lines(before)), 1)
        w.put("lv.csv", living([("10-02", "10:00:00", "타행이체", "신한 110222", 100_123, 0)]), folder="2026-10-28")
        before = w.csv()
        out, rid2 = self.ok_scan(w)
        rc, out = w.apply(rid2, *ACK)
        self.assertEqual(rc, 0, out)
        self.assertEqual(w.new_lines(before), [], out)

    def test_two_rounds_out_then_in(self):
        """출금 파일 먼저(상대 파일이 그 날짜를 안 덮음 → 기다림) → 입금 파일 다음 회차: 기다리는 이체를 먼저 소비, 행 0."""
        w = self.world(rows=BASE_ROWS)
        w.put("lv.csv", living([("10-02", "10:00:00", "타행이체", "신한 110222", 100_123, 0)]))
        out, rid = self.ok_scan(w)
        before = w.csv()
        rc, out = w.apply(rid, *ACK)
        self.assertEqual(rc, 0, out)
        self.assertEqual(w.new_lines(before), ['2026-10-02,이체,이체,신한은행,고정비 이체,"100,123",생활비,고정,'])
        self.assertEqual(w.ledger("SELECT count(*) FROM transfer WHERE in_key=''")[0][0], 1)
        w.put("sb.csv", shinhan_bank([("10-02", "10:00:01", "타행입금", "카카오 3333-01", 0, 100_123)]),
              folder="2026-10-28")
        before = w.csv()
        out, rid2 = self.ok_scan(w)
        rc, out = w.apply(rid2, *ACK)
        self.assertEqual(rc, 0, out)
        self.assertEqual(w.new_lines(before), [], out)
        self.assertEqual(w.ledger("SELECT count(*) FROM transfer WHERE in_key=''")[0][0], 0)

    def test_excluded_inbound_auto_and_ambiguous_name(self):
        """입금만 보이는 제외 계좌(커플통장 별명) → 자동 · 본인 이름(둘 이상 계좌) → 자동 0."""
        cfg = CONFIG + '\n[[transfer]]\naccount = "생활비"\ncounterpart = "커플통장"\ndirection = "in"\npatterns = ["커플정산"]\n'
        cfg = cfg.replace('patterns = ["커플모임통장"]', 'patterns = ["커플모임통장"]\ndirection = "out"')
        w = self.world(config=cfg, rows=BASE_ROWS)
        w.put("lv.csv", living([("10-03", "10:00:00", "입금", "커플정산", 0, 30_123),
                                ("10-04", "10:00:00", "입금", "가나다", 0, 40_000)]))
        out, rid = self.ok_scan(w)
        props = {p["금액"]: p for p in w.proposals(rid)}
        self.assertEqual(props["30123"]["상태"], "자동", out)
        self.assertEqual(props["30123"]["새 결제수단"], "커플통장")
        self.assertEqual(props["40000"]["상태"], "확인 필요")

    def test_same_day_same_amount_two_transfers(self):
        """같은 날 같은 금액 출금 2·입금 2 → 짝 2개, 행 2(이중 0)."""
        w = self.world(rows=BASE_ROWS)
        w.put("lv.csv", living([("10-05", "10:00:00", "타행이체", "신한 110222", 10_000, 0),
                                ("10-05", "10:01:00", "타행이체", "신한 110222", 10_000, 0)]))
        w.put("sb.csv", shinhan_bank([("10-05", "10:00:01", "타행입금", "카카오 3333-01", 0, 10_000),
                                      ("10-05", "10:01:01", "타행입금", "카카오 3333-01", 0, 10_000)]))
        out, rid = self.ok_scan(w)
        props = [p for p in w.proposals(rid) if p["금액"] == "10000"]
        self.assertEqual(sorted(p["상태"] for p in props), ["자동", "자동", "제외", "제외"], out)

    def test_memo_conflict(self):
        """시험 후보 8: 카드대금 적요가 이체 패턴에도 걸리면 '적요 충돌'(자동 0)."""
        cfg = CONFIG.replace("patterns = ['^현대카드$']", "patterns = ['^현대카드$', '^신한 110222$']")
        w = self.world(config=cfg, rows=BASE_ROWS)
        w.put("lv.csv", living([("10-02", "10:00:00", "타행이체", "신한 110222", 100_123, 0)]))
        out, rid = self.ok_scan(w)
        self.assertEqual([i["종류"] for i in w.items(rid)], ["memo_conflict"], out)


# ============================================================
# 카드대금 · 취소 · 할부 · 수입 (완료 기준 8·9·10·11)
# ============================================================

class CardTests(Base):
    def test_card_payment_no_row_and_reconcile(self):
        """카드사 출금은 행 0 · 청구 주기 카드 행 합계와 같으면 대조됨 · 다르면 apply 막힘(--ack-card)."""
        rows = BASE_ROWS
        w = self.world(rows=rows)
        w.put("h.csv", hyundai([("09-03", "10:00:00", "가상마트 2호점", 10013, "일시불", "10000021", "정상"),
                                ("09-20", "10:00:00", "가상마트 2호점", 5000, "일시불", "10000022", "정상"),
                                ("10-05", "10:00:00", "가상마트 2호점", 7013, "일시불", "10000023", "정상")],
                               period=("2026.09.01", "2026.10.27")))
        w.put("lv.csv", living([("10-12", "09:00:00", "현대카드", "현대카드", 15013, 0)]))
        out, rid = self.ok_scan(w)
        self.assertIn("카드대금 현대카드 2026-10: 차이 0 대조됨", out)
        props = [p for p in w.proposals(rid) if p["출처"] == "생활비"]
        self.assertEqual([p["상태"] for p in props], ["제외"])
        w.put("lv.csv", living([("10-12", "09:00:00", "현대카드", "현대카드", 16000, 0)]))
        out, rid = self.ok_scan(w)
        rc, out = w.apply(rid, confirm=False)
        self.assertEqual(rc, 1, out)
        self.assertIn("--ack-card 현대카드", out)
        rc, out = w.apply(rid, "--ack-card", "현대카드", "시험 사유", "--ack-balance", "생활비", "시험", confirm=False)
        self.assertEqual(rc, 0, out)

    def test_card_payment_missing_month(self):
        """카드 행이 있는 달인데 결제일을 덮는 통장 파일에 카드대금이 없으면 확인 필요."""
        w = self.world(rows=BASE_ROWS)
        w.put("h.csv", hyundai([("10-02", "10:00:00", "가상마트 2호점", 7013, "일시불", "10000023", "정상")]))
        w.put("lv.csv", living([("10-01", "09:00:00", "급여", "가상회사", 0, 3_123_457),
                                ("11-20", "09:00:00", "체크", "가상가게", 100, 0)]))
        out, rid = self.ok_scan(w, today=dt.date(2026, 11, 25))
        self.assertIn("cardpay_missing", {i["종류"] for i in w.items(rid)}, out)

    def test_rule_save_refused_for_card_names(self):
        """카드사명·'카드대금'·계좌명이 든 답은 규칙 저장 거부."""
        w = self.world()
        w.put("h.csv", hyundai([("10-06", "08:00:00", "새가게", 7013, "일시불", "10000003", "정상")]))
        out, rid = self.ok_scan(w)
        w.answer(rid, lambda d: d["종류"] == "new_merchant", 처리="추가", 대분류="식비", 소분류="외식",
                 세부내용="현대카드 연회비", **{"새 이름": "Y"})
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        self.assertIn("규칙 저장 거부", out)
        self.assertNotIn("새가게", open(os.path.join(w.rules, "분류규칙.csv"), encoding="utf-8-sig").read())

    def test_rule_save_allows_category_containing_account_word(self):
        """리허설 발견: 대분류 '주거생활비'처럼 계좌 짧은 이름이 들어 있는 분류는 저장된다('생활비' 그 자체만 거부)."""
        w = self.world()
        w.put("h.csv", hyundai([("10-06", "08:00:00", "새가게", 7013, "일시불", "10000003", "정상")]))
        out, rid = self.ok_scan(w)
        w.answer(rid, lambda d: d["종류"] == "new_merchant", 처리="추가", 대분류="주거생활비", 소분류="관리비",
                 세부내용="관리비", **{"새 이름": "Y"})
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        self.assertNotIn("규칙 저장 거부", out)
        self.assertIn("새가게", open(os.path.join(w.rules, "분류규칙.csv"), encoding="utf-8-sig").read())

    def test_cancel_then_reapproval_same_day_no_delete_question(self):
        """리허설 발견: 취소된 거래와 같은 날 같은 금액으로 다시 승인된 거래가 있고 손 행이 하나면, 손 행은 재승인 거래와 짝 —
        '손 행 삭제?' 질문을 내지 않는다(삭제로 답하면 실제 구매 행이 지워진다)."""
        rows = BASE_ROWS + ['2026-10-05,지출,식비,장보기,가상마트,"30,123",생활비,,현대카드']
        w = self.world(rows=rows)
        w.put("h.csv", hyundai([("10-05", "19:00:00", "가상마트 2호점", 30123, "일시불", "10000002", "취소"),
                                ("10-05", "19:05:00", "가상마트 2호점", 30123, "일시불", "10000009", "정상")]))
        out, rid = self.ok_scan(w)
        self.assertNotIn("cancel_row", [i["종류"] for i in w.items(rid)], out)
        st = sorted(p["상태"] for p in w.proposals(rid) if p["출처"] == "현대")
        self.assertEqual(st, ["제외", "짝"], out)

    def test_cancel_separate_row_neither_added(self):
        """취소 행(음수)과 원거래를 짝지어 둘 다 넣지 않는다."""
        w = self.world(rows=BASE_ROWS)
        w.put("s.csv", shinhan_card([("10-04", "가상마트 3호점", 48000, "20000001"),
                                     ("10-06", "가상마트 3호점", -48000, "20000001")]))
        out, rid = self.ok_scan(w)
        props = [p for p in w.proposals(rid) if p["출처"] == "신한카드"]
        self.assertEqual(sorted(p["상태"] for p in props), ["제외", "제외"], out)

    def test_cancel_in_later_round_asks_delete_then_not_again(self):
        """시험 후보 1: 취소가 둘째 파일에 옴 → 원거래가 CSV에 있으니 '삭제?' → 삭제 답 → 셋째 scan 질문 0."""
        w = self.world(rows=BASE_ROWS)
        w.put("s.csv", shinhan_card([("10-04", "가상마트 2호점", 48000, "20000001")]))
        out, rid = self.ok_scan(w)
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        w.put("s2.csv", shinhan_card([("10-04", "가상마트 2호점", 48000, "20000001"),
                                      ("10-06", "가상마트 2호점", -48000, "20000001")]), folder="2026-10-28")
        out, rid2 = self.ok_scan(w)
        self.assertEqual([i["종류"] for i in w.items(rid2)], ["cancel_row"], out)
        w.answer(rid2, lambda d: d["종류"] == "cancel_row", 처리="삭제")
        before = w.csv()
        rc, out = w.apply(rid2)
        self.assertEqual(rc, 0, out)
        self.assertNotIn("48,000", w.csv().decode("utf-8"))
        self.assertEqual(len(before.decode("utf-8-sig").split("\r\n")) - 1, len(w.csv().decode("utf-8-sig").split("\r\n")))
        out, rid3 = self.ok_scan(w)
        self.assertEqual(w.items(rid3), [], out)

    def test_status_change_cancel_detected(self):
        """원거래 행 자체가 '취소' 상태로 바뀐 파일(다음 회차) → '씀' 키도 건너뛰지 않고 확인 필요 · apply 막힘."""
        w = self.world(rows=BASE_ROWS)
        w.put("h.csv", hyundai([("10-05", "19:00:00", "가상마트 2호점", 30123, "일시불", "10000002", "정상")]))
        out, rid = self.ok_scan(w)
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        w.put("h.csv", hyundai([("10-05", "19:00:00", "가상마트 2호점", 30123, "일시불", "10000002", "취소")]))
        out, rid2 = self.ok_scan(w)
        self.assertEqual([i["종류"] for i in w.items(rid2)], ["cancel_row"], out)
        rc, out = w.apply(rid2)                      # 답 없음 → 보류(원장 값은 옛 값 그대로)
        self.assertEqual(rc, 0, out)
        out, rid3 = self.ok_scan(w)
        self.assertEqual([i["종류"] for i in w.items(rid3)], ["cancel_row"], out)

    def test_vanished_auto_row_blocks(self):
        """자동 행 거래가 파일 기간 안인데 파일에서 사라짐(행 삭제형 취소) → 취소 의심 · apply 막힘."""
        w = self.world(rows=BASE_ROWS)
        per = ("2026.10.04", "2026.10.27")
        w.put("h.csv", hyundai([("10-05", "19:00:00", "가상마트 2호점", 30123, "일시불", "10000002", "정상"),
                                ("10-06", "19:00:00", "가상마트 2호점", 1000, "일시불", "10000003", "정상")], period=per))
        out, rid = self.ok_scan(w)
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        w.put("h.csv", hyundai([("10-06", "19:00:00", "가상마트 2호점", 1000, "일시불", "10000003", "정상")], period=per))
        out, rid2 = self.ok_scan(w)
        self.assertEqual([i["종류"] for i in w.items(rid2)], ["vanished"], out)
        rc, out = w.apply(rid2)
        self.assertEqual(rc, 1, out)
        w.answer(rid2, lambda d: d["종류"] == "vanished", 처리="삭제")
        rc, out = w.apply(rid2)
        self.assertEqual(rc, 0, out)
        self.assertNotIn('"30,123"', w.csv().decode("utf-8"))

    def test_installment_schedule(self):
        """할부(총액 기준): 달마다 한 행·합 = 총액·미래 날짜 0·파일의 회차 행과 이중 0(시험 후보 2)."""
        rules = RULES + "가상전자,가전,전자제품,가상폰,,,자동,가상전자,2026-10-01,x\r\n"
        w = self.world(rows=BASE_ROWS, rules=rules)
        w.put("h.csv", hyundai([("10-18", "10:00:00", "가상전자", 1_000_013, "3개월", "10000031", "정상")],
                               period=("2026.10.01", "2026.10.27")))
        out, rid = self.ok_scan(w)
        self.assertEqual(read_bytes(os.path.join(w.run_dir(rid), f"초안_가계부_{rid}.csv")), w.csv(), out)
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        self.assertEqual(w.ledger("SELECT total, months, first_month FROM schedule"), [(1_000_013, 3, "2026-11")])
        # 12/20: 11/12·12/12 두 회차(333,337). 파일에 11월 청구 회차 행이 보여도 이중 0.
        w.put("h2.csv", hyundai([("11-12", "10:00:00", "가상전자", 333_337, "일시불", "10000032", "정상")],
                                period=("2026.10.28", "2026.12.20")), folder="2026-12-20")
        before = w.csv()
        out, rid2 = self.ok_scan(w, today=dt.date(2026, 12, 20))
        rc, out = w.apply(rid2, *NEW_MONTH, today=dt.date(2026, 12, 20))
        self.assertEqual(rc, 0, out)
        new = w.new_lines(before)
        self.assertEqual(new, ['2026-11-12,지출,가전,전자제품,가상폰,"333,337",생활비,,현대카드',
                               '2026-12-12,지출,가전,전자제품,가상폰,"333,337",생활비,,현대카드'], out)
        out, rid3 = self.ok_scan(w, today=dt.date(2027, 1, 20))
        rc, out = w.apply(rid3, *NEW_MONTH, today=dt.date(2027, 1, 20))
        self.assertEqual(rc, 0, out)
        self.assertIn('2027-01-12,지출,가전,전자제품,가상폰,"333,339",생활비,,현대카드', w.csv().decode("utf-8"))

    def test_installment_vs_hand_full_row(self):
        """할부 총액과 같은 손 행이 있으면 '전액 1행 vs 월별?' 확인 필요(자동 0)."""
        rows = BASE_ROWS + ['2026-10-18,지출,가전,전자제품,가상폰,"1,000,013",생활비,,현대카드']
        rules = RULES + "가상전자,가전,전자제품,가상폰,,,자동,가상전자,2026-10-01,x\r\n"
        w = self.world(rows=rows, rules=rules)
        w.put("h.csv", hyundai([("10-18", "10:00:00", "가상전자", 1_000_013, "3개월", "10000031", "정상")]))
        out, rid = self.ok_scan(w)
        self.assertEqual([i["종류"] for i in w.items(rid)], ["inst_vs_hand"], out)

    def test_answer_installment_for_new_merchant(self):
        """할부 표기가 있는 새 가맹점 — '할부'로 답하면 일정(이번 달은 회차 행 0)."""
        w = self.world(rows=BASE_ROWS)
        w.put("h.csv", hyundai([("10-18", "10:00:00", "가상폰가게", 900_000, "3개월", "10000041", "정상")]))
        out, rid = self.ok_scan(w)
        it = [i for i in w.items(rid) if i["종류"] == "new_merchant"]
        self.assertEqual(int(it[0]["할부 개월"]), 3, out)
        w.answer(rid, lambda d: d["종류"] == "new_merchant", 처리="할부", 대분류="식비", 소분류="외식",
                 세부내용="폰", **{"새 이름": "Y"})
        before = w.csv()
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        self.assertEqual(w.new_lines(before), [])
        self.assertEqual(w.ledger("SELECT months FROM schedule"), [(3,)])

    def test_income_rule_and_other_deposit(self):
        """완료 기준 11: 급여 규칙 → 자동 · 그 밖 입금 → '확인 필요(항목명)' 기본 기타 수익."""
        w = self.world(rows=BASE_ROWS)
        w.put("lv.csv", living([("10-09", "09:00:00", "입금", "아무개", 0, 12_000),
                                ("10-25", "09:00:00", "급여", "가상회사", 0, 3_123_457)]))
        out, rid = self.ok_scan(w)
        props = {p["금액"]: p for p in w.proposals(rid)}
        self.assertEqual((props["3123457"]["상태"], props["3123457"]["새 대분류"]), ("자동", "급여"), out)
        self.assertEqual((props["12000"]["상태"], props["12000"]["새 대분류"]), ("확인 필요", "기타 수익"))


# ============================================================
# 답 · 규칙 (완료 기준 12·13·14·20)
# ============================================================

class AnswerTests(Base):
    def test_rule_saved_then_auto_next_round(self):
        """'규칙 저장' 답 → 다음 회차 같은 가맹점(같은 원문)은 자동."""
        w = self.world(rows=BASE_ROWS)
        w.put("h.csv", hyundai([("10-06", "08:00:00", "새가게", 7013, "일시불", "10000003", "정상")]))
        out, rid = self.ok_scan(w)
        w.answer(rid, lambda d: d["종류"] == "new_merchant", 처리="추가", 대분류="식비", 소분류="외식", 세부내용="점심")
        before = w.csv()
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        self.assertEqual(w.new_lines(before), ['2026-10-06,지출,식비,외식,점심,"7,013",생활비,,현대카드'])
        w.put("h2.csv", hyundai([("10-09", "08:00:00", "새가게", 8013, "일시불", "10000013", "정상")]), folder="2026-10-28")
        out, rid2 = self.ok_scan(w)
        props = [p for p in w.proposals(rid2) if p["금액"] == "8013"]
        self.assertEqual(props[0]["상태"], "자동", out)

    def test_once_answer_no_rule(self):
        """'이번만' → 규칙 0, 다음 회차 다시 확인 필요."""
        w = self.world(rows=BASE_ROWS)
        w.put("h.csv", hyundai([("10-06", "08:00:00", "새가게", 7013, "일시불", "10000003", "정상")]))
        out, rid = self.ok_scan(w)
        w.answer(rid, lambda d: d["종류"] == "new_merchant", 처리="추가", 대분류="식비", 소분류="외식",
                 세부내용="점심", 규칙="이번만")
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        self.assertNotIn("새가게", open(os.path.join(w.rules, "분류규칙.csv"), encoding="utf-8-sig").read())

    def test_conflicting_answers_promote_every(self):
        """같은 키에 다른 답 → '매번' 승격 → 다음 회차 자동 0."""
        w = self.world(rows=BASE_ROWS)
        w.put("h.csv", hyundai([("10-06", "08:00:00", "새가게", 7013, "일시불", "10000003", "정상"),
                                ("10-08", "08:00:00", "새가게", 9013, "일시불", "10000004", "정상")]))
        out, rid = self.ok_scan(w)
        w.answer(rid, lambda d: d["금액"] == 7013, 처리="추가", 대분류="식비", 소분류="외식", 세부내용="점심")
        w.answer(rid, lambda d: d["금액"] == 9013, 처리="추가", 대분류="식비", 소분류="장보기", 세부내용="마트")
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        self.assertIn("매번", out)
        w.put("h2.csv", hyundai([("10-10", "08:00:00", "새가게", 1000, "일시불", "10000014", "정상")]), folder="2026-10-28")
        out, rid2 = self.ok_scan(w)
        self.assertEqual([p["상태"] for p in w.proposals(rid2) if p["금액"] == "1000"], ["확인 필요"], out)

    def test_platform_never_saved(self):
        """플랫폼 키 → 규칙 저장 불가·항상 확인 필요."""
        w = self.world(rows=BASE_ROWS)
        w.put("h.csv", hyundai([("10-07", "09:00:00", "가짜페이 결제", 20013, "일시불", "10000004", "정상")]))
        out, rid = self.ok_scan(w)
        w.answer(rid, lambda d: d["종류"] == "platform", 처리="추가", 대분류="식비", 소분류="외식", 세부내용="점심",
                 규칙="저장")
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        self.assertNotIn("가짜페이", open(os.path.join(w.rules, "분류규칙.csv"), encoding="utf-8-sig").read())

    def test_invalid_answer_reasked(self):
        """답 값 검증: 이름 목록 밖(새 이름 표시 없음)·계좌명 아닌 결제수단 → apply 막힘(되묻기)."""
        w = self.world(rows=BASE_ROWS)
        w.put("h.csv", hyundai([("10-06", "08:00:00", "새가게", 7013, "일시불", "10000003", "정상")]))
        out, rid = self.ok_scan(w)
        w.answer(rid, lambda d: d["종류"] == "new_merchant", 처리="추가", 대분류="없는분류", 소분류="외식", 세부내용="점심")
        rc, out = w.apply(rid)
        self.assertEqual(rc, 1, out)
        self.assertIn("CSV에 없는 이름", out)
        w.answer(rid, lambda d: d["종류"] == "new_merchant", 대분류="식비", 결제수단="현대카드")
        rc, out = w.apply(rid)
        self.assertEqual(rc, 1, out)
        self.assertIn("계좌명", out)

    def test_held_reappears_and_answers_inherited(self):
        """완료 기준 20: 답 안 한 항목은 보류로 남아 다음 회차(파일이 없어도) 다시 올라오고, 지난 답은 이어받는다."""
        w = self.world(rows=BASE_ROWS)
        path = w.put("h.csv", hyundai([("10-06", "08:00:00", "새가게", 7013, "일시불", "10000003", "정상"),
                                       ("10-07", "08:00:00", "다른가게", 3000, "일시불", "10000005", "정상")]))
        out, rid = self.ok_scan(w)
        rc, out = w.apply(rid)                       # 답 0 → 둘 다 보류
        self.assertEqual(rc, 0, out)
        os.remove(path)                              # 다음 회차엔 파일이 없다
        out, rid2 = self.ok_scan(w)
        self.assertEqual(len([i for i in w.items(rid2) if i["종류"] == "new_merchant"]), 2, out)
        w.answer(rid2, lambda d: d["금액"] == 7013, 처리="추가", 대분류="식비", 소분류="외식", 세부내용="점심")
        out, rid3 = self.ok_scan(w)                  # apply 없이 다시 scan → 답 이어받음
        inh = [i for i in w.items(rid3) if i["처리"] == "추가"]
        self.assertEqual(len(inh), 1, out)
        self.assertIn("이어받", inh[0]["메모"])

    def test_answer_command(self):
        """대화 답을 answer 명령으로 xlsx에 적는다(선택지 밖이면 멈춤)."""
        w = self.world(rows=BASE_ROWS)
        w.put("h.csv", hyundai([("10-06", "08:00:00", "새가게", 7013, "일시불", "10000003", "정상")]))
        out, rid = self.ok_scan(w)
        rc, out = w.run("answer", "--run", rid, "--item", "1", "--set", "처리=없는답")
        self.assertEqual(rc, 1, out)
        rc, out = w.run("answer", "--run", rid, "--item", "1", "--set", "처리=제외")
        self.assertEqual(rc, 0, out)
        self.assertEqual(w.items(rid)[0]["처리"], "제외")


# ============================================================
# 잔액 대조 (완료 기준 18 · 시험 후보 14)
# ============================================================

class BalanceTests(Base):
    def balance_world(self, start):
        w = self.world(rows=BASE_ROWS)
        w.put("lv.csv", living([("10-01", "09:00:00", "급여", "가상회사", 0, 3_123_457),
                                ("10-02", "09:00:00", "체크", "가상가게", 5_013, 0)], start=start))
        return w

    def computed_start(self):
        """CSV 사본으로 9/30 잔고를 계산해 파일 시작 잔액을 맞춘다(account_balance 공식)."""
        import pipeline
        w = self.world(rows=BASE_ROWS)
        with redirect_stdout(io.StringIO()):
            df = pipeline.load_csv(w.csv_path(), today=TODAY)
        # 9월 카드 행(12,345 + 48,123)은 CSV에선 9월에 빠졌지만 통장에선 10/12·10/15 결제일에 빠진다
        return pipeline.account_balance(df, "생활비 통장", until="2026-09-30")[0] + 60_468

    def test_formula_matches_account_balance(self):
        """파일 잔액 = CSV 잔고면 차액 0 통과 · 확인 필요(보류)는 항목으로 설명되어 잔차 0."""
        start = self.computed_start()
        w = self.balance_world(start)
        out, rid = self.ok_scan(w)
        self.assertIn("잔액 생활비: 통과(잔차 0) · 잔차 0", out)
        rep = open(os.path.join(w.run_dir(rid), "보고.md"), encoding="utf-8").read()
        self.assertIn("보류·확인 필요·제외(입금 + / 출금 −): -5,013", rep)

    def test_baseline_then_change_blocks(self):
        """첫 회차 잔차 ≠ 0 → --ack-balance 사유로 기준 차액 · 다음 회차 잔차가 그대로면 통과, 바뀌면 막힘."""
        start = self.computed_start() + 777
        w = self.balance_world(start)
        out, rid = self.ok_scan(w)
        rc, out = w.apply(rid, confirm=False)
        self.assertEqual(rc, 1, out)
        self.assertIn("--ack-balance 생활비", out)
        rc, out = w.apply(rid, "--ack-balance", "생활비", "시험 기준 차액")
        self.assertEqual(rc, 0, out)
        self.assertEqual(w.ledger("SELECT baseline FROM anchor WHERE account='생활비'"), [(777,)])
        out, rid2 = self.ok_scan(w)
        self.assertIn("통과(기준 차액 그대로)", out)
        # 손으로 CSV 행을 하나 더 적으면(파일에 없는 생활비 지출) — 내역에 없는 CSV 행 항으로 설명된다
        path = w.csv_path()
        data = w.csv().replace(b",,,,,,,,\r\n", "2026-10-02,지출,식비,외식,가상,1013,생활비,,\r\n,,,,,,,,\r\n".encode("utf-8"))
        with open(path, "wb") as f:
            f.write(data)
        out, rid3 = self.ok_scan(w)
        self.assertIn("통과(기준 차액 그대로)", out)
        # 원장 앵커: 파일을 지워도(파일 없는 계좌) 원장 거래로 다시 본다 — '생략' 아님
        shutil.rmtree(os.path.join(w.inbox, "2026-10-27"))
        out, rid4 = self.ok_scan(w)
        self.assertIn("잔액 생활비: 통과", out)


# ============================================================
# 공개 노출 · 자격 증명 (완료 기준 22 · 23 · 26)
# ============================================================

class ExposureTests(Base):
    def test_refuses_paths_inside_clone(self):
        """pc_import는 내역·rules·회차 폴더가 clone 안이면 시작하지 않는다."""
        P = pc_import.Paths(ROOT, os.path.dirname(ROOT), False)
        P.rules = os.path.join(ROOT, "rules")
        with self.assertRaises(pc_import.Stop):
            P.check_outside_clone()

    def test_no_binary_fixtures_in_repo(self):
        """저장소 tests/fixtures는 텍스트만(바이너리 내역 fixture 금지)."""
        fx = os.path.join(ROOT, "tests", "fixtures")
        for n in os.listdir(fx):
            data = read_bytes(os.path.join(fx, n))
            self.assertFalse(data[:4] in (b"PK\x03\x04", b"\xd0\xcf\x11\xe0") or data.startswith(b"SQLite"), n)
            self.assertFalse(n.lower().endswith((".xlsx", ".xls", ".sqlite", ".toml")), n)

    def test_outputs_have_no_secrets_and_report_no_merchants(self):
        """완료 기준 23 · 26: 회차 폴더·로그에 비밀번호·토큰 꼴 문자열 0 · 보고.md·log에 가맹점·개별 금액 0."""
        w = self.world()
        basic_files(w)
        out, rid = self.ok_scan(w)
        rc, out2 = w.apply(rid, "--ack-balance", "생활비", "시험", "--ack-balance", "신한은행", "시험")
        import re
        pat = re.compile(r"(?i)(password|passwd|비밀번호|token|ghp_[A-Za-z0-9]{10,}|github_pat_|secret|쿠키|cookie)")
        for r, _, fs in os.walk(w.run_dir(rid)):
            for n in fs:
                if n.endswith((".json", ".md", ".csv")):
                    self.assertIsNone(pat.search(open(os.path.join(r, n), encoding="utf-8-sig").read()), n)
        rep = open(os.path.join(w.run_dir(rid), "보고.md"), encoding="utf-8").read()
        for m in ("가상식당", "새가게", "가짜페이", "가상회사", "가나다", "7,013", "20,013"):
            self.assertNotIn(m, rep)
        if os.path.exists(os.path.join(w.home, "runs", "log.md")):
            log = open(os.path.join(w.home, "runs", "log.md"), encoding="utf-8").read()
            for m in ("가상", "새가게", "7,013"):
                self.assertNotIn(m, log)


# ============================================================
# 리허설 놀이터
# ============================================================

class RehearsalTests(Base):
    def test_rehearsal_writes_only_playground(self):
        """리허설: 운영 uploads·rules·work 쓰기 0 — CSV·rules·원장·회차는 work-rehearsal/내역/ 사본에만."""
        w = self.world()
        w.put("h.csv", hyundai([("10-05", "19:00:00", "가상마트 2호점", 30123, "일시불", "10000002", "정상")]))

        def tree():
            out = {}
            for root in (w.uploads, w.rules):
                for r, _, fs in os.walk(root):
                    for n in fs:
                        out[os.path.join(r, n)] = md5(read_bytes(os.path.join(r, n)))
            return out
        before = tree()
        rc, out = w.run("--rehearsal", "init")
        self.assertEqual(rc, 0, out)
        rc, out = w.run("--rehearsal", "scan")
        self.assertIn(rc, (0, 2), out)
        self.assertIn("[리허설]", out)
        runs = pc_import.run_dirs(pc_import.Paths(ROOT, w.home, True))
        rc, out = w.run("--rehearsal", "apply", "--run", runs[-1], "--confirm")
        self.assertEqual(rc, 0, out)
        self.assertEqual(before, tree())
        self.assertFalse(os.path.exists(os.path.join(w.home, "work")))
        play = os.path.join(w.home, "work-rehearsal", "내역", "uploads")
        self.assertEqual(len([n for n in os.listdir(play) if n.endswith(".csv")]), 1)


# ============================================================
# 자체 검토 막음 5건 회귀(2026-10-27) — 합성 값
# ============================================================

INST_RULES = RULES + "가상전자,가전,전자제품,가상폰,,,자동,가상전자,2026-10-01,x\r\n"


class ReviewFixTests(Base):
    def inst_world(self):
        w = self.world(rows=BASE_ROWS, rules=INST_RULES)
        w.put("h.csv", hyundai([("10-18", "10:00:00", "가상전자", 1_000_013, "3개월", "10000031", "정상")],
                               period=("2026.10.01", "2026.10.27")))
        out, rid = self.ok_scan(w)
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        return w

    def test_installment_cancel_ends_schedule(self):
        """막음 1: 할부 구매가 취소되면(별도 취소 행 · 상태 변경 둘 다) 일정이 끝나 다음 달부터 회차 행 0."""
        for variant in ("sep", "status"):
            with self.subTest(variant=variant):
                w = self.inst_world()
                rows = ([("10-18", "10:00:00", "가상전자", 1_000_013, "3개월", "10000031", "정상"),
                         ("10-28", "10:00:00", "가상전자", -1_000_013, "3개월", "10000031", "정상")]
                        if variant == "sep" else
                        [("10-18", "10:00:00", "가상전자", 1_000_013, "3개월", "10000031", "취소")])
                w.put("h2.csv", hyundai(rows, period=("2026.10.01", "2026.10.29")), folder="2026-10-29")
                out, rid = self.ok_scan(w, today=dt.date(2026, 10, 29))
                rc, out = w.apply(rid, today=dt.date(2026, 10, 29))
                self.assertEqual(rc, 0, out)
                self.assertEqual(w.ledger("SELECT mode FROM schedule"), [("cancelled",)])
                before = w.csv()
                out, rid = self.ok_scan(w, today=dt.date(2026, 11, 20))
                rc, out = w.apply(rid, *NEW_MONTH, today=dt.date(2026, 11, 20))
                self.assertEqual(rc, 0, out)
                self.assertEqual(w.new_lines(before), [], out)

    def test_cancelled_schedule_asks_written_rows(self):
        """막음 1: 이미 적힌 회차 행이 있는 일정이 취소되면 그 행마다 'CSV 행 삭제?'를 답할 때까지 묻는다."""
        w = self.inst_world()
        out, rid = self.ok_scan(w, today=dt.date(2026, 11, 13))
        rc, out = w.apply(rid, *NEW_MONTH, today=dt.date(2026, 11, 13))
        self.assertEqual(rc, 0, out)
        self.assertIn('"333,337"', w.csv().decode("utf-8"))
        w.put("h2.csv", hyundai([("10-18", "10:00:00", "가상전자", 1_000_013, "3개월", "10000031", "취소")],
                                period=("2026.10.01", "2026.11.14")), folder="2026-11-14")
        for _ in range(2):          # 답 없이 apply해도 다음 회차에 다시 묻는다
            out, rid = self.ok_scan(w, today=dt.date(2026, 11, 14))
            self.assertIn("cancel_row", [i["종류"] for i in w.items(rid)], out)
            rc, out = w.apply(rid, *NEW_MONTH, today=dt.date(2026, 11, 14))
            self.assertEqual(rc, 0, out)
        out, rid = self.ok_scan(w, today=dt.date(2026, 11, 14))
        w.answer(rid, lambda d: d["종류"] == "cancel_row", 처리="삭제")
        rc, out = w.apply(rid, *NEW_MONTH, today=dt.date(2026, 11, 14))
        self.assertEqual(rc, 0, out)
        self.assertNotIn('"333,337"', w.csv().decode("utf-8"))
        out, rid = self.ok_scan(w, today=dt.date(2026, 12, 20))
        self.assertNotIn("cancel_row", [i["종류"] for i in w.items(rid)], out)

    def test_row_deleted_cancel_with_old_folder_kept(self):
        """막음 2: 받은 날마다 새 폴더(옛 폴더 유지) — 최신 파일에서 행째 사라진 자동 행 거래는 취소 의심 · apply 막힘."""
        w = self.world(rows=BASE_ROWS)
        w.put("h.csv", hyundai([("10-05", "19:00:00", "가상마트 2호점", 30123, "일시불", "10000002", "정상"),
                                ("10-06", "19:00:00", "가상마트 2호점", 1013, "일시불", "10000003", "정상")],
                               period=("2026.10.04", "2026.10.27")))
        out, rid = self.ok_scan(w)
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        w.put("h.csv", hyundai([("10-06", "19:00:00", "가상마트 2호점", 1013, "일시불", "10000003", "정상")],
                               period=("2026.10.04", "2026.10.28")), folder="20261028")
        out, rid = self.ok_scan(w)
        self.assertEqual([i["종류"] for i in w.items(rid)], ["vanished"], out)
        rc, out = w.apply(rid)
        self.assertEqual(rc, 1, out)

    def test_stale_new_tx_not_auto_added(self):
        """막음 2 곁: 아직 적지 않은 거래가 최신 파일에서 사라졌으면 자동 행 0(취소 의심 확인 필요)."""
        w = self.world(rows=BASE_ROWS)
        w.put("h.csv", hyundai([("10-05", "19:00:00", "가상마트 2호점", 30123, "일시불", "10000002", "정상"),
                                ("10-06", "19:00:00", "가상마트 2호점", 1013, "일시불", "10000003", "정상")],
                               period=("2026.10.04", "2026.10.27")))
        w.put("h.csv", hyundai([("10-06", "19:00:00", "가상마트 2호점", 1013, "일시불", "10000003", "정상")],
                               period=("2026.10.04", "2026.10.28")), folder="20261028")
        out, rid = self.ok_scan(w)
        self.assertEqual([i["종류"] for i in w.items(rid)], ["vanished_new"], out)
        draft = read_bytes(os.path.join(w.run_dir(rid), f"초안_가계부_{rid}.csv")).decode("utf-8")
        self.assertNotIn('"30,123"', draft)

    def test_billing_row_in_later_round(self):
        """막음 4: 일정 회차 행이 적힌 다음 회차에 파일의 청구 회차 행이 처음 보여도 짝(이중 0)."""
        w = self.inst_world()
        out, rid = self.ok_scan(w, today=dt.date(2026, 11, 13))
        rc, out = w.apply(rid, *NEW_MONTH, today=dt.date(2026, 11, 13))
        self.assertEqual(rc, 0, out)
        w.put("h2.csv", hyundai([("11-10", "10:00:00", "가상전자", 333_337, "일시불", "10000032", "정상")],
                                period=("2026.10.28", "2026.11.20")), folder="2026-11-20")
        before = w.csv()
        out, rid = self.ok_scan(w, today=dt.date(2026, 11, 20))
        rc, out = w.apply(rid, *NEW_MONTH, today=dt.date(2026, 11, 20))
        self.assertEqual(rc, 0, out)
        self.assertEqual(w.new_lines(before), [], out)

    def test_billing_row_after_full_answer(self):
        """막음 4: 손 행 '전액' 답 뒤 청구 회차 행이 와도 자동 행 0."""
        rows = BASE_ROWS + ['2026-10-18,지출,가전,전자제품,가상폰,"1,000,013",생활비,,현대카드']
        w = self.world(rows=rows, rules=INST_RULES)
        w.put("h.csv", hyundai([("10-18", "10:00:00", "가상전자", 1_000_013, "3개월", "10000031", "정상")],
                               period=("2026.10.01", "2026.10.27")))
        out, rid = self.ok_scan(w)
        w.answer(rid, lambda d: d["종류"] == "inst_vs_hand", 처리="전액")
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        w.put("h2.csv", hyundai([("11-12", "10:00:00", "가상전자", 333_337, "일시불", "10000032", "정상")],
                                period=("2026.10.28", "2026.11.20")), folder="2026-11-20")
        before = w.csv()
        out, rid = self.ok_scan(w, today=dt.date(2026, 11, 20))
        rc, out = w.apply(rid, *NEW_MONTH, today=dt.date(2026, 11, 20))
        self.assertEqual(rc, 0, out)
        self.assertEqual(w.new_lines(before), [], out)

    def test_schedule_row_removed_from_csv_asks(self):
        """막음 5: 자동화가 적은 할부 회차 행이 CSV에서 사라지면 '사라진 행'으로 묻는다."""
        w = self.inst_world()
        out, rid = self.ok_scan(w, today=dt.date(2026, 11, 13))
        rc, out = w.apply(rid, *NEW_MONTH, today=dt.date(2026, 11, 13))
        self.assertEqual(rc, 0, out)
        path = w.csv_path()
        data = w.csv().replace('2026-11-12,지출,가전,전자제품,가상폰,"333,337",생활비,,현대카드\r\n'.encode("utf-8"), b"")
        with open(path, "wb") as f:
            f.write(data)
        out, rid = self.ok_scan(w, today=dt.date(2026, 11, 20))
        self.assertIn("gone_row", [i["종류"] for i in w.items(rid)], out)


# ============================================================
# 실물 형식 맞추기(2026-10-27 리허설) — 합성 값으로 같은 꼴
# ============================================================

FORMAT_CONFIG = r'''
start_date = "2026-10-01"
platform_keys = ["가짜페이"]

[[source]]
id = "현대H"
kind = "card"
card = "현대카드"
account = "생활비"
file_types = ["html"]
columns = ["승인일", "승인시각", "카드종류", "가맹점명", "승인금액", "할부개월", "승인번호", "승인구분"]
identity = ["-9876*"]
identity_column = "카드종류"
date_basis = "승인일"
date_formats = ["%Y년 %m월 %d일"]
time_formats = ["%H:%M"]
installment_mode = "total"
pay_day = 12
footer_re = "^-"
[source.roles]
date = "승인일"
time = "승인시각"
merchant = "가맹점명"
amount = "승인금액"
installment_months = "할부개월"
approval = "승인번호"
cancel = "승인구분"
[source.values]
"승인일" = { type = "s", re = '\d{4}년 \d{2}월 \d{2}일' }
"승인시각" = { type = "s", re = '\d{2}:\d{2}' }
"카드종류" = { type = "s" }
"가맹점명" = { type = "s" }
"승인금액" = { type = "s", re = '[\d,]+' }
"할부개월" = { type = "s", re = '\d+' }
"승인번호" = { type = "s", re = '\d{6}|\d{8}' }
"승인구분" = { type = "s", re = '전표접수|승인|취소|취소전표접수' }

[[source]]
id = "신한C"
kind = "card"
card = "신한카드"
account = "생활비"
file_types = ["xlsx", "text"]
columns = ["거래일", "카드구분", "이용카드", "가맹점명", "승인번호", "금액", "매입구분", "이용구분", "취소상태"]
identity = ["본인555"]
identity_column = "이용카드"
exclude_rows = [{ column = "카드구분", values = ["체크"] }]
date_basis = "승인일"
date_formats = ["%Y.%m.%d %H:%M"]
installment_mode = "total"
pay_day = 15
footer_re = '^(?!\d{4}\.)'
[source.roles]
date = "거래일"
merchant = "가맹점명"
amount = "금액"
approval = "승인번호"
cancel = "매입구분"
installment_months = "이용구분"
[source.values]
"거래일" = { type = "s", re = '\d{4}\.\d{2}\.\d{2} \d{2}:\d{2}' }
"카드구분" = { type = "s", re = '신용|체크' }
"이용카드" = { type = "s" }
"가맹점명" = { type = "s" }
"승인번호" = { type = "s", re = '\d{8}' }
"금액" = { type = ["n", "s"], re = '-?[\d,]+' }
"매입구분" = { type = "s", re = '결제확정|승인취소' }
"이용구분" = { type = "s", re = '일시불|할부\(\d+개월\)' }
"취소상태" = { type = "s", re = '취소', empty_ok = true }

[[source]]
id = "카뱅"
kind = "bank"
account = "생활비"
file_types = ["xlsx", "text"]
columns = ["", "거래일시", "구분", "거래금액", "거래 후 잔액", "거래구분", "내용"]
identity = ["***4321"]
date_formats = ["%Y.%m.%d %H:%M:%S"]
period_re = '(\d{4})\.(\d{2})\.(\d{2}) - (\d{4})\.(\d{2})\.(\d{2})'
direction_values = { "입금" = "in", "출금" = "out" }
[source.roles]
date = "거래일시"
merchant = "내용"
memo = ["거래구분", "내용"]
amount = "거래금액"
direction = "구분"
balance = "거래 후 잔액"
category = "거래구분"
[source.values]
"" = { type = "any", empty_ok = true }
"거래일시" = { type = "s", re = '\d{4}\.\d{2}\.\d{2} \d{2}:\d{2}:\d{2}' }
"구분" = { type = "s", re = '입금|출금' }
"거래금액" = { type = "s", re = '-?[\d,]+' }
"거래 후 잔액" = { type = "s", re = '[\d,]+' }
"거래구분" = { type = "s" }
"내용" = { type = "s", empty_ok = true }

[[source]]
id = "신한은행"
kind = "bank"
account = "신한은행"
file_types = ["text", "xls"]
columns = ["거래일자", "거래시간", "적요", "출금(원)", "입금(원)", "내용", "잔액(원)", "거래점"]
identity = ["110-555-666666"]
count_re = '총건수\s*(\d+)'
date_formats = ["%Y-%m-%d"]
[source.roles]
date = "거래일자"
time = "거래시간"
merchant = "내용"
memo = ["적요", "내용", "거래점"]
out = "출금(원)"
in = "입금(원)"
balance = "잔액(원)"
[source.values]
"거래일자" = { type = "s", re = '\d{4}-\d{2}-\d{2}' }
"거래시간" = { type = "s", re = '\d{2}:\d{2}:\d{2}' }
"적요" = { type = "s" }
"출금(원)" = { type = ["n", "s"], re = '[\d,]+' }
"입금(원)" = { type = ["n", "s"], re = '[\d,]+' }
"내용" = { type = "s", empty_ok = true }
"잔액(원)" = { type = ["n", "s"], re = '[\d,]+' }
"거래점" = { type = "s", empty_ok = true }

[[internal]]
account = "생활비"
name = "세이프박스"
field = "category"
patterns = ['^세이프박스$']

[[transfer]]
account = "생활비"
counterpart = "신한은행"
direction = "in"
field = "category"
patterns = ['^오픈뱅킹$']

[[transfer]]
account = "신한은행"
counterpart = "생활비"
direction = "out"
field = "memo"
patterns = ['^오픈뱅킹 이체$']

[[transfer_ambiguous]]
account = "생활비"
field = "merchant"
patterns = ['^홍길동$']
candidates = ["신한은행", "커플통장"]
default = "커플통장"
'''


def html_card(rows, card="4***-****-****-9876*"):
    """현대카드 실물 꼴 HTML(.xls) — 머리 위 제목 두 줄 · 스크립트 · 합계 행('-')."""
    def tr(cells):
        return "<tr>" + "".join(f"<td>{c}</td>" for c in cells) + "</tr>"
    body = [tr(["가상카드 실시간 이용내역"]), tr(["조회 결과"]),
            tr(["승인일", "승인시각", "카드종류", "가맹점명", "승인금액", "할부개월", "승인번호", "승인구분"])]
    for d, t, m, a, inst, ap, st in rows:
        body.append(tr([f"2026년 {d[:2]}월 {d[3:]}일", t, card, m, f"{a:,}", inst, ap, st]))
    body.append(tr(["-", "", "", "합계", "1", "", "", ""]))
    return ("<html><head><script>var x = '<tr><td>가짜</td></tr>';</script></head><body><table>"
            + "".join(body) + "</table></body></html>").encode("utf-8")


def kakao_like(rows, start=500_000, ident="****-**-***4321"):
    """카카오뱅크 실물 꼴: 머리 위 블록 · 머리 앞 빈 열 · 문자열 칸 · 부호 있는 금액 · 거래구분."""
    bal, data = start, []
    for d, t, gubun, kind, content, a in rows:
        bal += a
        data.append(["", f"2026.{d.replace('-', '.')} {t}", gubun, f"{a:,}", f"{bal:,}", kind, content])
    head = [["", "가상뱅크 거래내역"], [], ["", "성명", "홍길동", "", "조회기간", "2026.10.01 - 2026.10.27"],
            ["", "계좌번호", ident], []]
    return text_file(head, ["", "거래일시", "구분", "거래금액", "거래 후 잔액", "거래구분", "내용"], data)


def shinhan_like(rows, start=400_000, count=None, ident="110-555-666666"):
    bal, data = start, []
    for d, t, memo, o, i, content, branch in rows:
        bal = bal - o + i
        data.append([f"2026-{d}", t, memo, str(o), str(i), content, str(bal), branch])
    head = [["거래내역조회"], [], ["계좌번호", ident], ["조회기간", "2026.10.01 ~ 2026.10.27"],
            ["총건수", str(len(rows) if count is None else count)], []]
    return text_file(head, ["거래일자", "거래시간", "적요", "출금(원)", "입금(원)", "내용", "잔액(원)", "거래점"], data)


SC_COLS = ["거래일", "카드구분", "이용카드", "가맹점명", "승인번호", "금액", "매입구분", "이용구분", "취소상태"]


def shinhan_card_like(rows):
    data = [SC_COLS] + [[f"2026.{d.replace('-', '.')} {t}", kind, card, m, ap, a, buy, inst, st]
                        for d, t, kind, card, m, ap, a, buy, inst, st in rows]
    data.append(["", "", "", "", "", sum(r[6] for r in rows), "", "", ""])
    return lambda p: write_xlsx(p, data, sheet="카드이용내역")


class FormatTests(Base):
    def fworld(self, **kw):
        return self.world(config=FORMAT_CONFIG, **kw)

    def test_real_like_files_dispositions(self):
        """실물 꼴 4곳: HTML 카드 · 체크카드 행 제외 · 다른 승인번호 취소 · 할부 취소 · 세이프박스 · 오픈뱅킹 이체."""
        w = self.fworld(rows=BASE_ROWS)
        w.put("hyundaicard_20261027.xls", html_card([
            ("10-05", "19:00", "가상마트 2호점", 30123, "0", "10000002", "전표접수"),
            ("10-06", "09:05", "가상가게A", 7013, "0", "10000003", "승인"),
            ("10-07", "10:00", "가상가게B", 5013, "0", "10000004", "취소"),
        ]), folder="20261027")
        w.put("Shinhancard_20261027.xlsx", shinhan_card_like([
            ("10-04", "12:00", "신용", "본인555*", "가상서점", "20000001", 15013, "결제확정", "일시불", ""),
            ("10-06", "13:00", "신용", "본인555*", "가상서점", "20000002", -15013, "승인취소", "일시불", "취소"),
            ("10-08", "14:00", "체크", "본인777*", "가상주차", "20000003", 3013, "결제확정", "일시불", ""),
            ("10-18", "15:00", "신용", "본인555*", "가상폰가게", "20000004", 1_200_013, "결제확정", "할부(12개월)", ""),
            ("10-18", "15:30", "신용", "본인555*", "가상폰가게", "20000005", -1_200_013, "승인취소", "할부(12개월)", "취소"),
        ]), folder="20261027")
        w.put("카뱅_v.xlsx", kakao_like([
            ("10-02", "10:00:00", "출금", "세이프박스", "세이프박스", -100_013),
            ("10-03", "11:00:00", "입금", "오픈뱅킹", "홍길동", 50_013),
        ]), folder="20261027")
        w.put("신한은행.xls", shinhan_like([
            ("10-03", "10:59:00", "오픈뱅킹 이체", 50_013, 0, "가상송금", "자금부"),
        ]), folder="20261027")
        out, rid = self.ok_scan(w)
        self.assertIn("설정으로 뺀 행 1", out)
        props = {(p["출처"], p["금액"]): (p["상태"], p["사유"]) for p in w.proposals(rid)}
        self.assertEqual(props[("현대H", "30123")][0], "자동", out)           # 규칙 가상마트
        self.assertEqual(props[("현대H", "5013")][0], "제외")                 # 상태 '취소' 한 행
        self.assertEqual(props[("신한C", "15013")][0], "제외")                # 원거래 — 승인번호 다른 취소와 짝
        self.assertEqual(props[("신한C", "-15013")][0], "제외")
        self.assertEqual(props[("신한C", "1200013")][0], "제외")              # 취소된 할부 — 일정 0
        self.assertNotIn(("신한C", "3013"), props)                          # 체크카드 행은 설정으로 뺐다
        self.assertEqual(props[("카뱅", "100013")], ("제외", "내부 이동(세이프박스) — 같은 계좌 안의 칸 이동, 행 0"))
        self.assertEqual(props[("신한은행", "50013")][0], "자동")             # 이체 행은 출금 쪽 하나
        self.assertEqual(props[("카뱅", "50013")][0], "제외")                 # 오픈뱅킹 입금 = 그 이체의 짝
        draft = read_bytes(os.path.join(w.run_dir(rid), f"초안_가계부_{rid}.csv")).decode("utf-8")
        self.assertIn('2026-10-03,이체,이체,생활비,,"50,013",신한은행,,', draft)
        self.assertEqual(w.ledger("SELECT count(*) FROM schedule") if os.path.exists(
            os.path.join(w.rules, "원장.sqlite")) else [(0,)], [(0,)])
        rep = open(os.path.join(w.run_dir(rid), "보고.md"), encoding="utf-8").read()
        self.assertIn("세이프박스 칸 잔액(−)", rep)

    def test_count_and_direction_mismatch_fail(self):
        """총건수 ≠ 행 수 · 구분 열과 금액 부호가 다르면 그 출처 실패."""
        w = self.fworld(rows=BASE_ROWS)
        w.put("sb.xls", shinhan_like([("10-03", "10:59:00", "이체", 1013, 0, "가상", "자금부")], count=2))
        bad = kakao_like([("10-02", "10:00:00", "입금", "일반입금", "가상", -1013)])
        w.put("kb.xlsx", bad)
        out, rid = self.ok_scan(w)
        self.assertIn("총건수 2 ≠ 표의 행 1", out)
        self.assertIn("금액 부호", out)

    def test_posting_status_change_is_not_a_change(self):
        """현대카드 '승인' → '전표접수'(같은 승인번호)는 바뀜 질문 없이 끝남. '취소'로 바뀔 때만 묻는다."""
        w = self.fworld(rows=BASE_ROWS)
        w.put("h.xls", html_card([("10-05", "19:00", "가상마트 2호점", 30123, "0", "10000002", "승인")]),
              folder="20261027")
        out, rid = self.ok_scan(w)
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        w.put("h.xls", html_card([("10-05", "19:00", "가상마트 2호점", 30123, "0", "10000002", "전표접수")]),
              folder="20261028")
        out, rid = self.ok_scan(w)
        self.assertEqual(w.items(rid), [], out)
        w.put("h.xls", html_card([("10-05", "19:00", "가상마트 2호점", 30123, "0", "10000002", "취소")]),
              folder="20261029")
        out, rid = self.ok_scan(w)
        self.assertEqual([i["종류"] for i in w.items(rid)], ["cancel_row"], out)

    def test_transfer_vs_hand_income_row_asks(self):
        """리허설 발견: 옮긴 돈을 손으로 '받은 계좌 수입'으로 적어 온 행이 있으면 이체 행을 자동으로 더하지 않는다
        (받은 계좌에 두 번 잡힘) — 경로 다른 짝으로 묻는다. 손 행 날짜가 하루 달라도."""
        for hand_date in ("2026-10-03", "2026-10-04"):
            with self.subTest(hand_date=hand_date):
                rows = BASE_ROWS + [f'{hand_date},수입,기타 수익,,가상 정산,"50,013",생활비,,']
                w = self.fworld(rows=rows)
                w.put("kb.xlsx", kakao_like([("10-03", "11:00:00", "입금", "오픈뱅킹", "홍길동", 50_013)]))
                w.put("sb.xls", shinhan_like([("10-03", "10:59:00", "오픈뱅킹 이체", 50_013, 0, "가상송금", "자금부")]))
                out, rid = self.ok_scan(w)
                self.assertIn("pair_route", [i["종류"] for i in w.items(rid)], out)
                draft = read_bytes(os.path.join(w.run_dir(rid), f"초안_가계부_{rid}.csv"))
                self.assertEqual(draft, w.csv())

    def test_unique_transfer_rule_beats_own_name(self):
        """본인 이름 패턴(모호)과 유일한 이체 규칙(거래구분 오픈뱅킹)이 함께 맞으면 유일한 쪽."""
        w = self.fworld(rows=BASE_ROWS)
        w.put("kb.xlsx", kakao_like([("10-03", "11:00:00", "입금", "오픈뱅킹", "홍길동", 50_013),
                                     ("10-04", "11:00:00", "입금", "일반입금", "홍길동", 60_013)]))
        out, rid = self.ok_scan(w)
        kinds = {i["금액"]: i["종류"] for i in w.items(rid)}
        self.assertEqual(kinds.get(50013), "transfer_noout", out)   # 신한은행 파일 없음 → 출금 계좌 확인
        self.assertEqual(kinds.get(60013), "transfer_unknown")


# ============================================================
# 첫 실사용 전 고치기 C1~C3 — 카드값 이체 창 · 신한→생활비 이체 가르기 · 취소 짝 시각
# ============================================================

CARD_TRANSFER_CFG = CONFIG + r'''
[[transfer_ambiguous]]
account = "생활비"
patterns = ["라마바"]
candidates = ["커플통장", "신한은행"]
default = "커플통장"

[[card_transfer]]
card = "현대카드"
account = "생활비"
field = "merchant"
patterns = ['^가나다$']
days_before = 3

[[card_transfer]]
card = "신한카드"
account = "생활비"
field = "merchant"
patterns = ['^가나다$']
'''

FILL_CFG = FORMAT_CONFIG + r'''
[[transfer_pair]]
from = "신한은행"
to = "생활비"
"대분류" = "이체"
auto_only_with_income = true

[[income]]
account = "신한은행"
field = "merchant"
patterns = ['^가상급여처$']
"대분류" = "급여"
"세부내용" = "가상급여"
"고정여부" = "고정"
'''


class C123Tests(Base):
    # ---- C1 카드값 이체 창 ----
    def card_transfer_world(self):
        """BASE_ROWS 9월 카드 행: 현대카드 12,345 + 9,917(신한은행 귀속) = 22,262 · 신한카드 48,123.
        창: 현대 결제일 12 → 10/9~10/12 · 신한 결제일 15 → 10/12~10/15(기본 3일)."""
        w = self.world(config=CARD_TRANSFER_CFG, rows=BASE_ROWS)
        w.put("lv.csv", living([("10-10", "09:00:00", "이체", "가나다", 23_013, 0),     # 현대 창만
                                ("10-10", "10:00:00", "이체", "라마바", 12_013, 0),     # 창 안 · 다른 상대
                                ("10-12", "09:00:00", "이체", "가나다", 47_017, 0),     # 두 창 — 신한 합계에 가깝다
                                ("10-20", "09:00:00", "이체", "가나다", 30_013, 0)]))   # 창 밖
        return w

    def test_card_transfer_window_suggests_only(self):
        """C1a: 창 안 본인 이름 출금 → 확인 필요 · 제안 카드대금(카드명, 겹치면 직전 주기 합계에 가까운 카드) ·
        자동 행 0. 창 밖·다른 상대는 지금 동작(제안 커플통장) 그대로."""
        w = self.card_transfer_world()
        out, rid = self.ok_scan(w)
        items = {i["금액"]: i for i in w.items(rid)}
        self.assertEqual({k: v["종류"] for k, v in items.items()},
                         {23013: "transfer_unknown", 47017: "transfer_unknown", 30013: "transfer_unknown",
                          12013: "transfer_unknown"}, out)
        self.assertIn("제안 카드대금(현대카드)", items[23013]["사유"])
        self.assertIn("합계 22,262(CSV)", items[23013]["사유"])
        self.assertIn("제안 카드대금(신한카드)", items[47017]["사유"])
        self.assertIn("현대카드 결제일 2026-10-12", items[47017]["사유"])       # 후보·근거 둘 다
        self.assertIn("신한카드 결제일 2026-10-15", items[47017]["사유"])
        for amt in (30013, 12013):
            self.assertNotIn("카드대금", items[amt]["사유"])
            self.assertEqual(items[amt]["소분류"], "커플통장")                    # 제안 커플통장 그대로
        self.assertEqual({p["상태"] for p in w.proposals(rid) if p["출처"] == "생활비"}, {"확인 필요"})
        self.assertEqual(read_bytes(os.path.join(w.run_dir(rid), f"초안_가계부_{rid}.csv")), w.csv())
        # 카드대금 답은 그 카드의 [[card_transfer]] 패턴에 맞는 출금에만
        w.answer(rid, lambda d: d["금액"] == 12013, 처리="카드대금(현대카드)")
        rc, out = w.apply(rid, *ACK, confirm=False)
        self.assertEqual(rc, 1, out)
        self.assertIn("카드값 이체 설정", out)

    def test_card_transfer_answer_two_rounds(self):
        """C1b·C1c: 답 '카드대금' → 행 0 · 원장 card_transfer · 다음 회차 질문 0 · 카드대금 대조의 관측 출금 ·
        잔액 대조 '카드결제 계좌에 남은 돈' 항 · 같은 입력으로 재scan하면 기준 차액 그대로(ack 없이 apply)."""
        w = self.card_transfer_world()
        out, rid = self.ok_scan(w)
        resid1 = [x.split(" · ")[0] for x in read_report(w, rid) if x.startswith("  - 잔차")]
        self.assertEqual(len(resid1), 1, resid1)                                 # 생활비 하나(신한은행은 파일 없음)
        w.answer(rid, lambda d: d["금액"] == 23013, 처리="카드대금(현대카드)")
        w.answer(rid, lambda d: d["금액"] == 47017, 처리="카드대금(신한카드)")
        before = w.csv()
        rc, out = w.apply(rid, *ACK)
        self.assertEqual(rc, 0, out)
        self.assertEqual(w.new_lines(before), [], out)                           # 행 0
        self.assertIn("카드대금 현대카드 2026-10: 차이 23,013 카드결제 계좌 경유", out)   # 카드 파일 없음 — 파일 합계 0 대비
        self.assertIn("카드대금 신한카드 2026-10: 차이 47,017 카드결제 계좌 경유", out)
        extra = [json.loads(x) for x, in w.ledger("SELECT extra FROM state WHERE state='card_transfer'")]
        self.assertEqual(sorted((x["card"], x["pay_month"]) for x in extra),
                         [("신한카드", "2026-10"), ("현대카드", "2026-10")])
        base = w.ledger("SELECT baseline FROM anchor WHERE account='생활비'")[0][0]
        self.assertEqual(resid1, [f"  - 잔차 {base:,}"])                          # 답 전 scan = apply 계산
        # 둘째 회차(같은 입력·답): 질문 0 · 관측 출금으로 다시 쓰임 · 기준 차액 그대로 · ack 없이 apply
        out, rid2 = self.ok_scan(w)
        self.assertEqual(sorted(i["금액"] for i in w.items(rid2)), [12013, 30013], out)
        self.assertIn("카드대금 현대카드 2026-10: 차이 23,013 카드결제 계좌 경유", out)
        self.assertIn("잔액 생활비: 통과(기준 차액 그대로)", out)
        rep = read_report(w, rid2)
        self.assertIn("  - 카드결제 계좌에 남은 돈(−): 355", rep)                  # −((23,013−22,262)+(47,017−48,123))
        self.assertEqual([x.split(" · ")[0] for x in rep if x.startswith("  - 잔차")], resid1)
        rc, out = w.apply(rid2)
        self.assertEqual(rc, 0, out)

    def test_installment_month_residual_same_before_and_after_answer(self):
        """C1c(자체 검토): 카드값 이체 카드에 할부 회차 행이 있는 달도 답 전 scan 잔차 = apply 기준 차액 = 재scan
        (회차 행도 그 달 카드대금 이체가 관측돼야 결제됨)."""
        nov = dt.date(2026, 11, 27)
        w = self.world(config=CARD_TRANSFER_CFG, rows=BASE_ROWS)
        w.put("h.csv", hyundai([("10-05", "19:00:00", "가상마트 2호점", 30_039, "3개월", "10000101", "정상"),
                                ("10-06", "19:00:00", "가상마트 2호점", 5_013, "일시불", "10000102", "정상")],
                               period=("2026.10.01", "2026.11.27")))
        w.put("lv.csv", living([("10-10", "09:00:00", "이체", "가나다", 23_013, 0),
                                ("11-11", "09:00:00", "이체", "가나다", 15_013, 0),
                                ("11-20", "09:00:00", "체크", "가상가게", 1_013, 0)]))
        out, rid = self.ok_scan(w, today=nov)
        resid1 = [x.split(" · ")[0] for x in read_report(w, rid) if x.startswith("  - 잔차")]
        w.answer(rid, lambda d: d["종류"] == "transfer_unknown" and d["금액"] in (23013, 15013),
                 처리="카드대금(현대카드)")
        rc, out = w.apply(rid, *ACK, *NEW_MONTH, today=nov)
        self.assertEqual(rc, 0, out)
        self.assertIn('2026-11-12,지출,식비,장보기,가상마트,"10,013",생활비,,현대카드', w.csv().decode("utf-8"))
        base = w.ledger("SELECT baseline FROM anchor WHERE account='생활비'")[0][0]
        self.assertEqual(resid1, [f"  - 잔차 {base:,}"], out)
        out, rid2 = self.ok_scan(w, today=nov)
        self.assertIn("잔액 생활비: 통과(기준 차액 그대로)", out)

    def test_card_transfer_counts_as_observed_payment(self):
        """C1b: 카드대금 이체 답은 관측된 카드대금 — '카드 행이 있는 달인데 카드대금 출금이 없다' 질문이 없어진다."""
        w = self.world(config=CARD_TRANSFER_CFG, rows=BASE_ROWS)
        w.put("h.csv", hyundai([("10-06", "08:00:00", "가상마트 2호점", 7_013, "일시불", "10000061", "정상")]))
        w.put("lv.csv", living([("10-01", "09:00:00", "급여", "가상회사", 0, 3_123_457),
                                ("11-11", "09:00:00", "이체", "가나다", 7_013, 0),
                                ("11-20", "09:00:00", "체크", "가상가게", 1_013, 0)]))
        today = dt.date(2026, 11, 25)
        out, rid = self.ok_scan(w, today=today)
        self.assertIn("cardpay_missing", [i["종류"] for i in w.items(rid)], out)
        w.answer(rid, lambda d: d["종류"] == "transfer_unknown" and d["금액"] == 7013, 처리="카드대금(현대카드)")
        out, rid2 = self.ok_scan(w, today=today)                                 # 답 이어받음
        self.assertNotIn("cardpay_missing", [i["종류"] for i in w.items(rid2)], out)
        self.assertIn("카드대금 현대카드 2026-11", out)

    # ---- C2 신한→생활비 이체 가르기 ----
    def fill_world(self):
        rows = BASE_ROWS + ['2026-10-05,지출,문화생활비,어플/멤버쉽,가상구독,"6,017",신한은행,고정,현대카드']
        w = self.world(config=FILL_CFG, rows=rows)
        w.put("sb.xls", shinhan_like([
            ("10-07", "09:00:00", "급여", 0, 1_234_017, "가상급여처", ""),
            ("10-07", "09:30:00", "오픈뱅킹 이체", 1_234_017, 0, "가상송금", "자금부"),
            ("10-09", "10:00:00", "오픈뱅킹 이체", 6_017, 0, "가상송금", "자금부"),
        ]))
        w.put("kb.xlsx", kakao_like([("10-07", "09:31:00", "입금", "오픈뱅킹", "홍길동", 1_234_017),
                                     ("10-09", "10:01:00", "입금", "오픈뱅킹", "홍길동", 6_017)]))
        return w

    def test_fill_transfer_needs_same_day_income(self):
        """C2a: 같은 날 같은 금액 자동 수입이 있는 쌍만 자동 이체 행 · 그 밖은 확인 필요(제안 제외 · 귀속 카드 행 사유) ·
        자동 제외 0."""
        w = self.fill_world()
        out, rid = self.ok_scan(w)
        props = {(p["출처"], p["금액"], p["방향"]): p for p in w.proposals(rid)}
        self.assertEqual(props[("신한은행", "1234017", "in")]["상태"], "자동", out)            # 수입 규칙
        self.assertEqual(props[("신한은행", "1234017", "out")]["상태"], "자동")                 # 이체 행
        self.assertIn("자동 수입", props[("신한은행", "1234017", "out")]["사유"])
        self.assertEqual(props[("카뱅", "1234017", "in")]["상태"], "제외")                      # 이체 짝
        self.assertEqual(props[("신한은행", "6017", "out")]["상태"], "확인 필요")
        self.assertEqual(props[("카뱅", "6017", "in")]["상태"], "제외")
        self.assertIn("답을 따른다", props[("카뱅", "6017", "in")]["사유"])
        it = [i for i in w.items(rid) if i["금액"] == 6017]
        self.assertEqual([i["종류"] for i in it], ["transfer_fill"], out)
        self.assertIn("제안 제외", it[0]["사유"])
        self.assertIn("CSV 16행(2026-10-05 현대카드)", it[0]["사유"])
        draft = read_bytes(os.path.join(w.run_dir(rid), f"초안_가계부_{rid}.csv")).decode("utf-8")
        self.assertIn('2026-10-07,이체,이체,생활비,,"1,234,017",신한은행,,', draft)
        self.assertIn('2026-10-07,수입,급여,,가상급여,"1,234,017",신한은행,고정,', draft)
        self.assertNotIn('"6,017",신한은행,,', draft)

    def test_fill_answers_two_rounds(self):
        """C2b: 답 '제외' → 행 0 · 다음 회차 질문 0 / 답 '이체' → 이체 한 행(양쪽 키) · 다음 회차 새 행 0."""
        for act in ("제외", "이체"):
            with self.subTest(act=act):
                w = self.fill_world()
                out, rid = self.ok_scan(w)
                w.answer(rid, lambda d: d["종류"] == "transfer_fill", 처리=act)
                before = w.csv()
                rc, out = w.apply(rid, *ACK)
                self.assertEqual(rc, 0, out)
                fill = [x for x in w.new_lines(before) if '"6,017"' in x]
                if act == "제외":
                    self.assertEqual(fill, [], out)
                    self.assertEqual(sorted(s for s, in w.ledger(
                        "SELECT state FROM state WHERE key LIKE '%/6017/%'")), ["suppressed", "suppressed"])
                else:
                    self.assertEqual(fill, ['2026-10-09,이체,이체,생활비,,"6,017",신한은행,,'], out)
                    tr = w.ledger("SELECT out_key, in_key FROM transfer WHERE amount=6017")
                    self.assertEqual(len(tr), 1)
                    self.assertTrue(tr[0][0] and tr[0][1], tr)
                before = w.csv()
                out, rid2 = self.ok_scan(w)
                self.assertEqual([i for i in w.items(rid2) if i["금액"] == 6017], [], out)
                rc, out = w.apply(rid2, *ACK)
                self.assertEqual(rc, 0, out)
                self.assertEqual(w.new_lines(before), [], out)

    # ---- C3 취소 짝 시각 ----
    def test_cancel_before_purchase_time_is_orphan(self):
        """C3a: 같은 날 취소 먼저·결제 나중 → 짝 0 · 짝 없는 취소는 확인 필요 · 결제는 할부 일정
        (실행일 뒤 회차 행 0). 대조군: 결제 먼저·취소 나중은 지금처럼 짝(둘 다 행 0)."""
        w = self.world(config=FORMAT_CONFIG, rows=BASE_ROWS)
        w.put("sc.xlsx", shinhan_card_like([
            ("10-23", "13:07", "신용", "본인555*", "가상폰가게", "20000011", -900_017, "승인취소", "할부(6개월)", "취소"),
            ("10-23", "14:52", "신용", "본인555*", "가상폰가게", "20000012", 900_017, "결제확정", "할부(6개월)", ""),
        ]))
        out, rid = self.ok_scan(w)
        kinds = {i["금액"]: i for i in w.items(rid)}
        self.assertEqual(kinds[-900017]["종류"], "cancel_orphan", out)
        self.assertEqual(kinds[900017]["종류"], "new_merchant")
        self.assertEqual(int(kinds[900017]["할부 개월"]), 6)
        w.answer(rid, lambda d: d["금액"] == 900017, 처리="할부", 대분류="식비", 소분류="외식", 세부내용="폰",
                 **{"새 이름": "Y"})
        before = w.csv()
        rc, out = w.apply(rid, *ACK)
        self.assertEqual(rc, 0, out)
        self.assertEqual(w.new_lines(before), [], out)                           # 11/15부터 — 실행일 뒤
        self.assertEqual(w.ledger("SELECT total, months, first_month FROM schedule"), [(900_017, 6, "2026-11")])
        w2 = self.world(config=FORMAT_CONFIG, rows=BASE_ROWS)
        w2.put("sc.xlsx", shinhan_card_like([
            ("10-23", "13:07", "신용", "본인555*", "가상폰가게", "20000012", 900_017, "결제확정", "할부(6개월)", ""),
            ("10-23", "14:52", "신용", "본인555*", "가상폰가게", "20000011", -900_017, "승인취소", "할부(6개월)", "취소"),
        ]))
        out, rid = self.ok_scan(w2)
        self.assertEqual(sorted(p["상태"] for p in w2.proposals(rid) if p["출처"] == "신한C"), ["제외", "제외"], out)
        self.assertEqual(w2.items(rid), [], out)

    def test_cancel_with_later_repurchase_does_not_pick_earlier_purchase(self):
        """C3 검토 막음: 취소 시각 뒤 같은 날 재결제가 있으면, 남은 앞선 같은 구매 하나와 자동으로 짝짓지 않는다
        (그 구매가 조용히 빠지지 않게) — 묻는다. 승인번호가 같은 거래가 취소 뒤면 가맹점·금액 짝으로 넘어가지 않는다."""
        for cancel_ap, kind in (("20000032", "cancel_ambiguous"), ("20000033", "cancel_orphan")):
            with self.subTest(kind=kind):
                w = self.world(config=FORMAT_CONFIG, rows=BASE_ROWS)
                w.put("sc.xlsx", shinhan_card_like([
                    ("10-04", "12:10", "신용", "본인555*", "가상꽃집", "20000031", 20_017, "결제확정", "일시불", ""),
                    ("10-21", "13:10", "신용", "본인555*", "가상꽃집", cancel_ap, -20_017, "승인취소", "일시불", "취소"),
                    ("10-21", "13:40", "신용", "본인555*", "가상꽃집", "20000033", 20_017, "결제확정", "일시불", ""),
                ]))
                out, rid = self.ok_scan(w)
                kinds = {(i["날짜"], i["금액"]): i["종류"] for i in w.items(rid)}
                self.assertEqual(kinds[("2026-10-21", -20017)], kind, out)
                props = {(p["날짜"], p["금액"]): p["상태"] for p in w.proposals(rid) if p["출처"] == "신한C"}
                self.assertEqual(props[("2026-10-04", "20017")], "확인 필요", out)      # 앞선 구매 — 조용히 제외 0
                self.assertEqual(props[("2026-10-21", "20017")], "확인 필요")

# ============================================================
# 규칙 칸 '매번'(이번 완료 기준 E1~E5 · 기존 12·13·14·20) — 합성 값만
# ============================================================

EVERY_RULES = RULES + "새가게,식비,외식,점심,,,매번,새가게,2026-10-01,20261001-000000\r\n"
A_VALS = {"대분류": "식비", "소분류": "외식", "세부내용": "점심"}
B_VALS = {"대분류": "식비", "소분류": "장보기", "세부내용": "마트"}


class EveryTests(Base):
    def rule_lines(self, w, key):
        import csv as _csv
        with open(os.path.join(w.rules, "분류규칙.csv"), encoding="utf-8-sig", newline="") as f:
            return [r for r in _csv.DictReader(f) if r["키"] == key]

    def rules_text(self, w):
        with open(os.path.join(w.rules, "분류규칙.csv"), encoding="utf-8-sig") as f:
            return f.read()

    def one(self, w, rid, amount):
        its = [i for i in w.items(rid) if i["금액"] == amount]
        self.assertEqual(len(its), 1, its)
        return its[0]

    def vals(self, it):
        return {c: it[c] for c in ("대분류", "소분류", "세부내용")}

    def test_every_answer_asks_next_round_with_suggestion(self):
        """E1: 새 가맹점 답 규칙 '매번' → 규칙 줄(물음 매번) 하나 → 다음 회차 ask_every · 제안 = 그 답 · 자동 행 0."""
        w = self.world(rows=BASE_ROWS)
        r1 = ("10-06", "08:00:00", "새가게", 6113, "일시불", "10000003", "정상")
        w.put("h.csv", hyundai([r1]))
        out, rid = self.ok_scan(w)
        it = self.one(w, rid, 6113)
        self.assertEqual((it["종류"], it["규칙"]), ("new_merchant", "저장"), out)
        w.answer(rid, lambda d: d["금액"] == 6113, 처리="추가", 규칙="매번", **A_VALS)
        before = w.csv()
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        self.assertEqual(w.new_lines(before), ['2026-10-06,지출,식비,외식,점심,"6,113",생활비,,현대카드'])
        self.assertEqual([(r["물음"], r["소분류"], r["원문"]) for r in self.rule_lines(w, "새가게")],
                         [("매번", "외식", "새가게")])
        # 다음 회차 파일은 기간이 겹치게 받는다(공백 0)
        w.put("h2.csv", hyundai([r1, ("10-09", "08:00:00", "새가게", 7213, "일시불", "10000013", "정상")]),
              folder="2026-10-28")
        out, rid2 = self.ok_scan(w)
        it = self.one(w, rid2, 7213)
        self.assertEqual((it["종류"], it["규칙"], it["결제수단"], it["비고"]), ("ask_every", "매번", "생활비", "현대카드"),
                         out)
        self.assertEqual(self.vals(it), A_VALS)
        self.assertEqual([p["상태"] for p in w.proposals(rid2) if p["금액"] == "7213"], ["확인 필요"])
        before = w.csv()
        rc, out = w.apply(rid2)                       # 답 0 → 보류 · 자동 행 0
        self.assertEqual(rc, 0, out)
        self.assertEqual(w.new_lines(before), [])

    def test_every_suggestion_follows_changed_answer(self):
        """E2: ask_every 답이 마지막 줄과 다르면 줄 +1 → 셋째 회차 제안 = 새 답. 같으면 줄 +0 · '이번만'이면 줄 +0."""
        w = self.world(rows=BASE_ROWS)
        rows = [("10-06", "08:00:00", "새가게", 6113, "일시불", "10000003", "정상")]
        w.put("h.csv", hyundai(rows))
        out, rid = self.ok_scan(w)
        w.answer(rid, lambda d: d["금액"] == 6113, 처리="추가", 규칙="매번", **A_VALS)
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        # 둘째 회차: 제안 A → B로 답(규칙 칸은 기본 '매번' 그대로)
        rows.append(("10-09", "08:00:00", "새가게", 7213, "일시불", "10000013", "정상"))
        w.put("h2.csv", hyundai(rows), folder="2026-10-28")
        out, rid2 = self.ok_scan(w)
        self.assertEqual(self.vals(self.one(w, rid2, 7213)), A_VALS, out)
        w.answer(rid2, lambda d: d["금액"] == 7213, 처리="추가", **B_VALS)
        rc, out = w.apply(rid2)
        self.assertEqual(rc, 0, out)
        self.assertEqual([(r["물음"], r["소분류"]) for r in self.rule_lines(w, "새가게")],
                         [("매번", "외식"), ("매번", "장보기")])
        # 셋째 회차: 제안 = B. 하나는 제안 그대로(같은 답 → 줄 +0), 하나는 다른 값을 '이번만'(줄 +0)
        rows += [("10-11", "08:00:00", "새가게", 8313, "일시불", "10000023", "정상"),
                 ("10-12", "08:00:00", "새가게", 9413, "일시불", "10000024", "정상")]
        w.put("h3.csv", hyundai(rows), folder="2026-10-29")
        out, rid3 = self.ok_scan(w)
        self.assertEqual(self.vals(self.one(w, rid3, 8313)), B_VALS, out)
        self.assertEqual(self.vals(self.one(w, rid3, 9413)), B_VALS)
        w.answer(rid3, lambda d: d["금액"] == 8313, 처리="추가")
        w.answer(rid3, lambda d: d["금액"] == 9413, 처리="추가", 규칙="이번만", 대분류="교통비", 소분류="대중교통",
                 세부내용="버스")
        before = w.csv()
        rc, out = w.apply(rid3)
        self.assertEqual(rc, 0, out)
        self.assertEqual(w.new_lines(before), ['2026-10-11,지출,식비,장보기,마트,"8,313",생활비,,현대카드',
                                               '2026-10-12,지출,교통비,대중교통,버스,"9,413",생활비,,현대카드'])
        self.assertEqual(len(self.rule_lines(w, "새가게")), 2, out)
        rows.append(("10-14", "08:00:00", "새가게", 5513, "일시불", "10000033", "정상"))
        w.put("h4.csv", hyundai(rows), folder="2026-10-30")
        out, rid4 = self.ok_scan(w)
        it = self.one(w, rid4, 5513)
        self.assertEqual((it["종류"], self.vals(it)), ("ask_every", B_VALS), out)

    def test_every_same_round_last_answer_wins_and_save_stays_every(self):
        """M3: 같은 회차 같은 키 답 둘 → 날짜·항목 순서로 마지막 답이 다음 제안. ask_every에 '저장' 답 → '매번'으로 남고 보고 한 줄."""
        w = self.world(rows=BASE_ROWS, rules=EVERY_RULES)
        w.put("h.csv", hyundai([("10-06", "08:00:00", "새가게", 6113, "일시불", "10000003", "정상"),
                                ("10-09", "08:00:00", "새가게", 7213, "일시불", "10000013", "정상")]))
        out, rid = self.ok_scan(w)
        self.assertEqual({(i["종류"], i["규칙"]) for i in w.items(rid)}, {("ask_every", "매번")}, out)
        w.answer(rid, lambda d: d["금액"] == 6113, 처리="추가", 규칙="저장", **B_VALS)   # 앞 날짜: B
        w.answer(rid, lambda d: d["금액"] == 7213, 처리="추가", **A_VALS)               # 뒤 날짜: A(지난 줄과 같음)
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        self.assertIn("저장 답도 매번으로 남음", out)
        self.assertEqual([(r["물음"], r["소분류"]) for r in self.rule_lines(w, "새가게")],
                         [("매번", "외식"), ("매번", "장보기"), ("매번", "외식")])
        w.put("h2.csv", hyundai([("10-12", "08:00:00", "새가게", 8313, "일시불", "10000023", "정상")]), folder="2026-10-28")
        out, rid2 = self.ok_scan(w)
        it = self.one(w, rid2, 8313)
        self.assertEqual((it["종류"], self.vals(it)), ("ask_every", A_VALS), out)
        self.assertEqual([p["상태"] for p in w.proposals(rid2) if p["금액"] == "8313"], ["확인 필요"])

    def test_rule_column_values_checked(self):
        """E3 · M1: 규칙 선택지 저장·매번·이번만 · answer --set 규칙=그밖 → 멈춤 · xlsx의 그 밖 값 → 답 오류(되묻기)로 apply 막힘."""
        import openpyxl
        w = self.world(rows=BASE_ROWS)
        w.put("h.csv", hyundai([("10-06", "08:00:00", "새가게", 6113, "일시불", "10000003", "정상")]))
        out, rid = self.ok_scan(w)
        wb = openpyxl.load_workbook(os.path.join(w.run_dir(rid), "확인필요.xlsx"))
        cols = {c[0].value: [x.value for x in c[1:] if x.value is not None] for c in wb["목록"].iter_cols()}
        self.assertEqual(cols["규칙"], ["저장", "매번", "이번만"])
        rc, out = w.run("answer", "--run", rid, "--item", "1", "--set", "처리=추가", "--set", "규칙=자동")
        self.assertEqual(rc, 1, out)
        self.assertEqual(w.items(rid)[0]["처리"], None)              # 멈추면 쓰기 0
        rc, out = w.run("answer", "--run", rid, "--item", "1", "--set", "규칙=매번")
        self.assertEqual(rc, 0, out)
        self.assertEqual(w.items(rid)[0]["규칙"], "매번")
        w.answer(rid, lambda d: d["금액"] == 6113, 처리="추가", 규칙="자동", **A_VALS)
        before = w.csv()
        rc, out = w.apply(rid)
        self.assertEqual(rc, 1, out)
        self.assertIn("답 오류(되묻기)", out)
        self.assertIn("규칙 '자동'", out)
        self.assertEqual(w.csv(), before)
        w.answer(rid, lambda d: d["금액"] == 6113, 규칙="이번만")
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        self.assertEqual(self.rule_lines(w, "새가게"), [])

    def test_every_answer_follows_refusals(self):
        """E4: '매번' 답도 거부 목록(카드사명·카드대금·계좌명·플랫폼 키·지출 아님)을 따른다 — ask_every 답도 같다."""
        w = self.world(rows=BASE_ROWS, rules=EVERY_RULES)
        w.put("h.csv", hyundai([("10-06", "08:00:00", "다른가게", 6113, "일시불", "10000003", "정상"),
                                ("10-07", "09:00:00", "가짜페이 결제", 7213, "일시불", "10000004", "정상"),
                                ("10-08", "08:00:00", "세번째가게", 8313, "일시불", "10000005", "정상"),
                                ("10-09", "08:00:00", "새가게", 9413, "일시불", "10000006", "정상"),
                                ("10-11", "08:00:00", "네번째가게", 5513, "일시불", "10000007", "정상")]))
        out, rid = self.ok_scan(w)
        kinds = {i["금액"]: i["종류"] for i in w.items(rid)}
        self.assertEqual(kinds, {6113: "new_merchant", 7213: "platform", 8313: "new_merchant", 9413: "ask_every",
                                 5513: "new_merchant"}, out)
        new = {"새 이름": "Y"}
        w.answer(rid, lambda d: d["금액"] == 6113, 처리="추가", 규칙="매번", 대분류="식비", 소분류="외식",
                 세부내용="현대카드 연회비", **new)                                              # 카드사명
        w.answer(rid, lambda d: d["금액"] == 7213, 처리="추가", 규칙="매번", **A_VALS)          # 플랫폼 키
        w.answer(rid, lambda d: d["금액"] == 8313, 처리="추가", 규칙="매번", 대분류="식비", 소분류="외식",
                 세부내용="커플통장 정산", **new)                                                # 계좌명
        w.answer(rid, lambda d: d["금액"] == 9413, 처리="추가", 대분류="식비", 소분류="외식",
                 세부내용="신한카드 카드대금", **new)                                            # ask_every · 카드대금
        w.answer(rid, lambda d: d["금액"] == 5513, 처리="추가", 규칙="매번", 구분="수입", **A_VALS)  # 지출 아님
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        self.assertEqual(out.count("규칙 저장 거부"), 3, out)
        text = self.rules_text(w)
        for m in ("다른가게", "가짜페이", "세번째가게", "네번째가게"):
            self.assertNotIn(m, text)
        self.assertEqual(len(self.rule_lines(w, "새가게")), 1)
        self.assertIn("규칙 0줄", out)

    def test_every_answer_refused_for_deposit(self):
        """E4 입금: 규칙 저장 가능 종류(손 행 짝 질문)의 입금 거래에 '매번' 답 → 규칙 0."""
        w = self.world(rows=BASE_ROWS)
        w.put("lv.csv", living([("10-04", "09:00:00", "입금", "아무개", 0, 3_123_457)]))
        out, rid = self.ok_scan(w)
        its = [i for i in w.items(rid) if i["금액"] == 3_123_457]
        self.assertEqual(len(its), 1, out)
        self.assertTrue(pc_import.KINDS[its[0]["종류"]][2], its[0]["종류"])
        w.answer(rid, lambda d: d["금액"] == 3_123_457, 처리="추가", 규칙="매번", 대분류="급여", 세부내용="10월 급여")
        rc, out = w.apply(rid, *ACK)
        self.assertEqual(rc, 0, out)
        self.assertNotIn("아무개", self.rules_text(w))

    def test_installment_answer_rule_lines(self):
        """E5 · M4: '할부' 답도 규칙 칸을 따른다 — 저장 → 다음 회차 자동 일정 · 매번 → 다음 회차 ask_every(제안 = 답)."""
        w = self.world(rows=BASE_ROWS)
        rows = [("10-18", "10:00:00", "가상폰가게", 900_000, "3개월", "10000041", "정상"),
                ("10-19", "10:00:00", "가상패드가게", 600_013, "3개월", "10000042", "정상")]
        w.put("h.csv", hyundai(rows))
        out, rid = self.ok_scan(w)
        self.assertEqual({i["종류"] for i in w.items(rid)}, {"new_merchant"}, out)
        w.answer(rid, lambda d: d["금액"] == 900_000, 처리="할부", 대분류="식비", 소분류="외식", 세부내용="폰",
                 **{"새 이름": "Y"})
        w.answer(rid, lambda d: d["금액"] == 600_013, 처리="할부", 규칙="매번", 대분류="식비", 소분류="외식",
                 세부내용="패드", **{"새 이름": "Y"})
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        self.assertEqual([r["물음"] for r in self.rule_lines(w, "가상폰가게")], ["자동"])
        self.assertEqual([r["물음"] for r in self.rule_lines(w, "가상패드가게")], ["매번"])
        rows += [("10-21", "10:00:00", "가상폰가게", 301_217, "3개월", "10000051", "정상"),
                 ("10-21", "11:00:00", "가상패드가게", 330_013, "3개월", "10000052", "정상")]
        w.put("h2.csv", hyundai(rows), folder="2026-10-28")       # 기간이 겹치게(공백 0)
        out, rid2 = self.ok_scan(w)
        props = {p["금액"]: p["상태"] for p in w.proposals(rid2)}
        self.assertEqual(props["301217"], "제외", out)                   # 규칙 자동 → 할부 일정(행은 다음 달부터)
        it = self.one(w, rid2, 330_013)
        self.assertEqual((it["종류"], it["세부내용"]), ("ask_every", "패드"), out)
        w.answer(rid2, lambda d: d["금액"] == 330_013, 처리="할부", **{"새 이름": "Y"})   # 제안 그대로 → 줄 +0
        rc, out = w.apply(rid2)
        self.assertEqual(rc, 0, out)
        self.assertEqual(len(self.rule_lines(w, "가상패드가게")), 1, out)
        self.assertEqual(w.ledger("SELECT count(*) FROM schedule"), [(4,)])

    def test_installment_ask_every_on_auto_key_promotes(self):
        """자체 검토 막음: 할부 거래는 옛 키가 '자동'이어도(처음 보는 원문) ask_every로 묻는다 — 제안 그대로 '매번'·'저장'
        답이면 '매번'으로 승격되고(줄 0이 아님), 다음 회차 같은 키의 본 원문 거래도 확인 필요(자동 0)."""
        for rule in ("매번", "저장"):
            with self.subTest(rule=rule):
                w = self.world(rows=BASE_ROWS)
                rows = [("10-18", "10:00:00", "가상마트 3호점", 450_013, "3개월", "10000061", "정상")]
                w.put("h.csv", hyundai(rows))
                out, rid = self.ok_scan(w)
                it = self.one(w, rid, 450_013)
                self.assertEqual((it["종류"], it["규칙"]), ("ask_every", "매번"), out)
                w.answer(rid, lambda d: d["금액"] == 450_013, 처리="할부", 규칙=rule, **{"새 이름": "Y"})
                rc, out = w.apply(rid)
                self.assertEqual(rc, 0, out)
                self.assertIn("'매번'으로 승격", out)
                self.assertNotIn("저장 답도 매번", out)
                self.assertEqual([r["물음"] for r in self.rule_lines(w, "가상마트")], ["자동", "매번"])
                rows.append(("10-21", "10:00:00", "가상마트 2호점", 31_013, "일시불", "10000062", "정상"))
                w.put("h2.csv", hyundai(rows), folder="2026-10-28")
                out, rid2 = self.ok_scan(w)
                self.assertEqual([p["상태"] for p in w.proposals(rid2) if p["금액"] == "31013"], ["확인 필요"], out)
                self.assertEqual(self.one(w, rid2, 31_013)["종류"], "ask_every")

    def test_every_key_same_value_any_kind_adds_nothing(self):
        """임의 결정: 이미 '매번'이 된 키에 마지막 줄과 같은 값이면 종류·규칙 칸과 상관없이 줄 0(같은 회차 '매번' 뒤 '저장')."""
        w = self.world(rows=BASE_ROWS)
        w.put("h.csv", hyundai([("10-06", "08:00:00", "새가게", 6113, "일시불", "10000003", "정상"),
                                ("10-09", "08:00:00", "새가게", 7213, "일시불", "10000013", "정상")]))
        out, rid = self.ok_scan(w)
        w.answer(rid, lambda d: d["금액"] == 6113, 처리="추가", 규칙="매번", **A_VALS)
        w.answer(rid, lambda d: d["금액"] == 7213, 처리="추가", 규칙="저장", **A_VALS)
        rc, out = w.apply(rid)
        self.assertEqual(rc, 0, out)
        self.assertEqual([r["물음"] for r in self.rule_lines(w, "새가게")], ["매번"])


# ============================================================
# E1 — 통장 밖 칸(세이프박스)은 칸 이름마다 한 번만 뺀다(첫 실사용 오류 E1 · 완료 기준 18)
# ============================================================

BOX_CFG = FORMAT_CONFIG + r'''
[[internal]]
account = "생활비"
name = "세이프박스"
field = "merchant"
patterns = ['^세이프박스$']

[[internal_balance]]
account = "생활비"
name = "세이프박스"
date = "2026-10-01"
amount = 200013
'''

TWO_BOX_CFG = BOX_CFG + r'''
[[internal]]
account = "생활비"
name = "모임칸"
field = "merchant"
patterns = ['^모임칸$']
'''

BOX_MOVES = [("10-02", "10:00:00", "출금", "세이프박스", "세이프박스", -100_013),   # 거래구분·내용 두 줄 다 맞음
             ("10-04", "10:00:00", "입금", "일반입금", "세이프박스", 30_013)]       # 내용 줄만 맞음


def residual(out, acct="생활비"):
    m = re.search(rf"잔액 {acct}: [^\n]* · 잔차 (-?[\d,]+)", out)
    return int(m.group(1).replace(",", ""))


class E1Tests(Base):
    """같은 칸을 적요 칸마다 [[internal]] 여러 줄로 두면 칸 잔액·이동액이 줄마다 빠졌다(잔차가 칸 잔액만큼 어긋남)."""

    EDGE = [("10-01", "09:00:00", "입금", "이자", "가상이자", 2_113),     # 두 판의 파일 시작·끝 날짜를 같게
            ("10-06", "09:00:00", "입금", "이자", "가상이자", 3_117)]

    def box_scan(self, w, config, moves, start):
        with open(os.path.join(w.rules, "계좌.toml"), "w", encoding="utf-8") as f:
            f.write(config)
        w.put("카뱅_v.xlsx", kakao_like(sorted(self.EDGE + moves), start=start), folder="20261027")
        out, rid = self.ok_scan(w)
        return out, rid, [x for x in read_report(w, rid) if "칸 잔액" in x]

    def test_same_box_two_lines_subtracts_once(self):
        """세이프박스 두 줄(거래구분 · 내용) → 칸 잔액 항 한 줄 = −(시작 잔액 + 순이동) · 잔차는 칸 없는 판과 같다.
        파일을 지운 다음 회차(원장 앵커 — 원장 거래를 설정 줄에 다시 맞춤)도 한 번만."""
        w = self.world(config=FORMAT_CONFIG, rows=BASE_ROWS)
        out0, _, lines0 = self.box_scan(w, FORMAT_CONFIG, [], 500_000)
        self.assertEqual(lines0, ["  - 세이프박스 칸 잔액(−) — 시작 잔액 미입력(기준 차액에 남음): 0"])
        out, rid, lines = self.box_scan(w, BOX_CFG, BOX_MOVES, 500_000 - 200_013)   # 칸에 든 돈은 파일 잔액 밖
        props = {p["금액"]: p["상태"] for p in w.proposals(rid)}
        self.assertEqual((props["100013"], props["30013"]), ("제외", "제외"))
        self.assertEqual(lines, ["  - 세이프박스 칸 잔액(−): -270,013"])
        self.assertEqual(residual(out), residual(out0), out)
        rc, out = w.apply(rid, "--ack-balance", "생활비", "시험 기준 차액")
        self.assertEqual(rc, 0, out)
        shutil.rmtree(os.path.join(w.inbox, "20261027"))
        out2, rid2 = self.ok_scan(w)
        self.assertIn("잔액 생활비: 통과(기준 차액 그대로)", out2)
        self.assertEqual([x for x in read_report(w, rid2) if "칸 잔액" in x], lines)

    def test_two_boxes_count_own_moves(self):
        """같은 계좌의 다른 칸(모임칸)은 제 이동만 — 시작 잔액이 없는 칸은 as_of 뒤 이동만 뺀다."""
        w = self.world(config=FORMAT_CONFIG, rows=BASE_ROWS)
        out0, _, _ = self.box_scan(w, FORMAT_CONFIG, [], 500_000)
        out, rid, lines = self.box_scan(w, TWO_BOX_CFG, BOX_MOVES + [
            ("10-05", "10:00:00", "출금", "일반출금", "모임칸", -7_013)], 500_000 - 200_013)
        self.assertEqual(lines, ["  - 세이프박스 칸 잔액(−): -270,013",
                                 "  - 모임칸 칸 잔액(−) — 시작 잔액 미입력(기준 차액에 남음): -7,013"])
        self.assertEqual(residual(out), residual(out0), out)


def read_report(w, rid):
    """회차 보고.md 줄 목록."""
    with open(os.path.join(w.run_dir(rid), "보고.md"), encoding="utf-8") as f:
        return f.read().split("\n")


class ExposeMoreTests(unittest.TestCase):
    def test_blob_hits_catch_common_shapes(self):
        """막음 3: '(주)…' 상호 · 시각 뺀 날짜+금액 · h:mm · MM-DD(시험 도우미 꼴) · 가계부 꼴 행도 잡는다."""
        sys.path.insert(0, os.path.join(ROOT, "scripts"))
        import pc_run
        import datetime as _dt
        v = {"lines": set(), "ids": set(), "stamps": set(), "words": set()}
        pc_run._row_values([("s", "2026.10.05 19:00:05"), ("s", "(주)가상유통 2호점"), ("s", "30,123")], v)
        pc_run._row_values([("d", _dt.datetime(2026, 10, 6, 9, 5)), ("s", "가상식당"), ("n", 17013.0)], v)
        v["words"] |= {"(주)가상유통 2호점"}
        cases = {
            "A": 'hyundai([("10-05", "19:00:05", "x", 30123, "일시불", "1", "정상")])',
            "C": '2026.10.05,(주)가상유통 2호점,"30,123"',
            "D": "2026-10-06 9:05,x,\"17,013\"",
            "F": '2026-10-06,지출,식비,외식,가상식당,"17,013",생활비,,현대카드',
            "G": "M = '(주)가상유통 2호점'",
        }
        for k, line in cases.items():
            self.assertTrue(pc_run.blob_hits(line.encode("utf-8"), v), k)
        self.assertEqual(pc_run.blob_hits('2026-10-07,x,"30,124"\n(주)다른상점'.encode("utf-8"), v), [])
        # origin/main에 이미 있던 줄은 새 노출이 아니다 · 날짜 꼴 8자리는 식별 숫자가 아니다
        line = '2026-10-06,지출,식비,외식,가상식당,"17,013",생활비,,현대카드'
        self.assertEqual(pc_run.blob_hits(line.encode("utf-8"), v, frozenset({line})), [])
        self.assertEqual(pc_run._ids_in("2026-10-01 · 20261001"), set())
        self.assertEqual(pc_run._ids_in("계좌 110-555-666666"), {"110555666666"})

    def test_real_values_cover_hyundai_cells_and_small_amounts(self):
        """검증 B1: 현대카드 꼴 날짜 셀('2026년 10월 09일')·100원 미만 금액(이자)도 (날짜, 금액) 대조에 든다 —
        시각 뺀 꼴·월-일 꼴·현대카드 셀 꼴로 옮겨도 잡히고, 다른 날짜·금액은 통과."""
        sys.path.insert(0, os.path.join(ROOT, "scripts"))
        import pc_run
        home = tempfile.mkdtemp(prefix="pcexpo-")
        self.addCleanup(rmtree, home)
        inbox = os.path.join(home, "uploads", "내역", "2026-10-27")
        os.makedirs(inbox)
        os.makedirs(os.path.join(home, "rules"))
        with open(os.path.join(home, "rules", "계좌.toml"), "w", encoding="utf-8") as f:
            f.write(FORMAT_CONFIG)
        files = {"h.xls": html_card([("10-09", "13:47", "가상상점", 31_417, "0", "40000001", "승인")]),
                 "kb.xlsx": kakao_like([("10-11", "09:00:00", "입금", "이자", "이자", 7)])}
        for name, data in files.items():
            with open(os.path.join(inbox, name), "wb") as f:
                f.write(data)
        v = pc_run.real_values(pc_run.Ctx(ROOT, home, False), None)
        for line in ('ROW = ("2026-10-09", 31417)', "| 2026-10-09 13:47 | 31,417 |", '("10-09", 31417)',
                     '"2026년 10월 09일","31,417"', "10월 9일 31,417원", 'ROW = ("2026-10-11", 7)'):
            self.assertTrue(pc_run.blob_hits(line.encode("utf-8"), v), line)
        for line in ('ROW = ("2026-10-09", 31418)', '"2026년 10월 10일","31,417"', 'ROW = ("2026-10-11", 8)'):
            self.assertEqual(pc_run.blob_hits(line.encode("utf-8"), v), [], line)


if __name__ == "__main__":
    unittest.main()
